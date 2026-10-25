"""Material boundary tests: no real cloud calls and no personal source fixtures."""
import io
from decimal import Decimal

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import db, materials
from app.integrations import ProviderError, tencent_headers
from app.main import create_app


@pytest.fixture
def material_app(tmp_path, monkeypatch):
    monkeypatch.setenv("CLOUD_ANALYSIS_ENABLED", "false")
    monkeypatch.setattr(materials, "configuration", lambda: {"enabled": False, "secret_id": "", "secret_key": "", "deepseek_key": "", "deepseek_model": "test-no-network", "general_limit": 10, "seal_limit": 10, "max_analysis_pages": 12})
    application = create_app(storage_dir=tmp_path / "storage", database_url=f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    clients = []

    def login(username):
        client = TestClient(application)
        response = client.post("/api/login", json={"username": username, "password": "Demo@2026"})
        assert response.status_code == 200, response.text
        client.headers["X-CSRF-Token"] = response.json()["csrf_token"]
        clients.append(client)
        return client

    yield application, login
    for client in clients:
        client.close()


def png(width=300, height=200, color="white"):
    image = Image.new("RGB", (width, height), color)
    stream = io.BytesIO()
    image.save(stream, "PNG")
    return stream.getvalue()


def create_claim(client):
    bootstrap = client.get("/api/bootstrap").json()
    year = (bootstrap.get("active_year") or bootstrap.get("current_year") or bootstrap["years"][0])["id"]
    category = next(c["id"] for c in bootstrap["categories"] if c["code"] == "COLLECTIVE")
    response = client.post("/api/applications", json={"year_id": year, "category_id": category})
    assert response.status_code == 200, response.text
    claim = response.json()
    return claim, year, bootstrap["classes"][0]["id"]


def upload(client, claim_id):
    response = client.post(f"/api/applications/{claim_id}/files", files={"file": ("proof.png", png(), "image/png")})
    assert response.status_code == 200, response.text
    return response.json()


def test_image_signatures_coordinate_bounds_and_tc3():
    with pytest.raises(HTTPException):
        materials.render_upload(b"<script>not an image</script>", "fake.png")
    with pytest.raises(HTTPException):
        materials.render_upload(png(), "fake.pdf")
    with pytest.raises(HTTPException):
        materials.checked_box([0, 0, float("nan"), 4], 100, 100)
    with pytest.raises(HTTPException):
        materials.checked_box([99, 0, 10, 10], 100, 100)
    assert materials.checked_box([1.2, 2.1, 3.3, 4.1], 100, 100) == [1, 2, 4, 5]
    config = {"secret_id": "EXAMPLE-ID", "secret_key": "EXAMPLE-KEY", "region": "ap-shanghai"}
    a = tencent_headers("GeneralAccurateOCR", b'{"x":1}', config, 1551113065)
    b = tencent_headers("GeneralAccurateOCR", b'{"x":2}', config, 1551113065)
    assert "/2019-02-25/ocr/tc3_request" in a["Authorization"]
    assert a["Authorization"] != b["Authorization"]
    assert a["X-TC-Action"] == "GeneralAccurateOCR"
    assert a["X-TC-Version"] == "2018-11-19"


def test_upload_private_access_is_scoped_and_admin_has_no_materials(material_app):
    _, login = material_app
    student = login("student")
    claim, _, _ = create_claim(student)
    record = upload(student, claim["id"])
    private_url = record["url"]
    assert student.get(private_url).status_code == 200
    assert login("student2").get(private_url).status_code == 403
    assert login("reviewer1").get(private_url).status_code == 200
    assert login("class2_reviewer1").get(private_url).status_code == 403
    assert login("admin").get(private_url).status_code == 403
    assert login("counselor").get(record["pages"][0]["url"]).status_code == 200
    response = student.post(f"/api/applications/{claim['id']}/files", files={"file": ("bad.png", b"not an image", "image/png")})
    assert response.status_code == 400


def test_public_crop_is_blocked_without_ocr_and_numeric_rows_are_redacted(material_app):
    _, login = material_app
    student = login("student")
    claim, year, class_id = create_claim(student)
    record = upload(student, claim["id"])
    page_id = record["pages"][0]["id"]
    body = {"page_id": page_id, "evidence_type": "PARTICIPATION", "box": [0, 0, 200, 120]}
    blocked = student.post(f"/api/applications/{claim['id']}/crops", json=body)
    assert blocked.status_code == 200, blocked.text
    assert blocked.json()["public_asset_status"] == "BLOCKED"
    assert blocked.json()["public_url"] is None
    with db.write() as conn:
        run = db.add(conn, "recognition_runs", page_id=page_id, provider="TENCENT", api_name="GeneralAccurateOCR", parameters_hash="test-only", status="SUCCEEDED", result_json={}, started_at=db.now(), finished_at=db.now())
        db.add(conn, "recognition_regions", run_id=run["id"], region_no=0, kind="TEXT", text="学号 DEMO202401004", confidence=Decimal("1"), polygon_json=[[20, 20], [160, 20], [160, 45], [20, 45]], extra_json={})
    ready = student.post(f"/api/applications/{claim['id']}/crops", json=body)
    assert ready.status_code == 200, ready.text
    assert ready.json()["public_asset_status"] == "READY"
    with db.read() as conn:
        crop = db.one(conn, "evidence_crops", ready.json()["id"])
        public_image = Image.open(materials.asset_path(crop["public_asset_uri"]))
        assert public_image.size == (300, 50)  # Whole original width, one quarter of original height.
        assert public_image.getpixel((40, 30)) == (37, 50, 68)
        assert public_image.getpixel((250, 40)) == (255, 255, 255)
    # Even a fully prepared image has no public access before a completed review and open publicity.
    assert login("student2").get(ready.json()["public_url"]).status_code == 403
    updated = student.put(f"/api/applications/{claim['id']}", json={"title": "测试集体活动", "requested_score": "0.30", "payload_json": {}})
    assert updated.status_code == 200, updated.text
    submitted = student.post(f"/api/applications/{claim['id']}/submit", json={})
    assert submitted.status_code == 200, submitted.text
    reviewer = login("reviewer1")
    decision = reviewer.post(f"/api/applications/{claim['id']}/review", json={"action": "APPROVE"})
    assert decision.status_code == 200, decision.text
    counselor = login("counselor")
    opened = counselor.put("/api/publicity", json={"class_id": class_id, "year_id": year, "is_open": True})
    assert opened.status_code == 200, opened.text
    mate = login("student2")
    assert mate.get(ready.json()["public_url"]).status_code == 200
    assert mate.get(record["url"]).status_code == 403
    assert login("class2_reviewer1").get(ready.json()["public_url"]).status_code == 403


def test_crop_cannot_point_at_another_claim_file_and_late_analysis_is_rejected(material_app):
    _, login = material_app
    student = login("student")
    first, _, _ = create_claim(student)
    second, _, _ = create_claim(student)
    record = upload(student, first["id"])
    wrong = student.post(f"/api/applications/{second['id']}/crops", json={"page_id": record["pages"][0]["id"], "evidence_type": "VALIDITY", "validity_kind": "NAME", "box": [0, 0, 80, 80]})
    assert wrong.status_code == 400
    with db.write() as conn:
        app = db.one(conn, "applications", first["id"])
        revision = db.one(conn, "application_revisions", app["working_revision_id"])
        context = materials._app_context(conn, first["id"])
        analysis = db.add(conn, "ai_analyses", application_id=first["id"], input_revision_id=revision["id"], input_fingerprint=materials.revision_fingerprint(conn, revision), score_rule_version_id=context["class_year"]["score_rule_version_id"], model_name="test-no-network", prompt_version="test", status="SUCCEEDED", result_json={"suggestion": {"title": "old suggestion", "requested_score": "0.30", "payload_json": {}}})
    change_response = student.put(f"/api/applications/{first['id']}", json={"title": "学生已自行修订", "requested_score": "0.60", "payload_json": {}})
    assert change_response.status_code == 200
    applied = student.post(f"/api/applications/{first['id']}/apply-analysis", json={"analysis_id": analysis["id"]})
    assert applied.status_code == 409, applied.text
    assert student.get(f"/api/applications/{first['id']}").json()["revision"]["title"] == "学生已自行修订"


def test_ocr_budget_reserves_failures_and_cache_avoids_second_request(material_app, monkeypatch):
    _, login = material_app
    student = login("student")
    claim, _, _ = create_claim(student)
    record = upload(student, claim["id"])
    with db.read() as conn:
        page = db.one(conn, "file_pages", record["pages"][0]["id"])
    calls = []

    def fake_success(action, image, config):
        calls.append(action)
        return {"RequestId": "test-no-network", "TextDetections": []}

    monkeypatch.setattr(materials, "tencent_ocr", fake_success)
    config = {"general_limit": 1, "seal_limit": 1}
    first = materials._ocr_page(page, "GeneralAccurateOCR", config)
    second = materials._ocr_page(page, "GeneralAccurateOCR", config)
    assert first["id"] == second["id"]
    assert calls == ["GeneralAccurateOCR"]

    def fake_failure(action, image, config):
        calls.append(action)
        raise ProviderError("ResourceUnavailable.ResourcePackageRunOut", "已耗尽")

    monkeypatch.setattr(materials, "tencent_ocr", fake_failure)
    with pytest.raises(ProviderError):
        materials._ocr_page(page, "SealOCR", config)
    with pytest.raises(ProviderError, match="本月本地 OCR"):
        materials._ocr_page(page, "SealOCR", config)
    assert calls.count("SealOCR") == 1


def test_first_crop_delete_after_rejection_preserves_the_frozen_evidence(material_app):
    _, login = material_app
    student = login("student")
    claim, _, _ = create_claim(student)
    record = upload(student, claim["id"])
    crop_response = student.post(f"/api/applications/{claim['id']}/crops", json={"page_id": record["pages"][0]["id"], "evidence_type": "VALIDITY", "validity_kind": "NAME", "box": [0, 0, 100, 100]})
    assert crop_response.status_code == 200
    crop_id = crop_response.json()["id"]
    saved = student.put(f"/api/applications/{claim['id']}", json={"title": "待修改证明", "requested_score": "0.30", "payload_json": {}})
    assert saved.status_code == 200
    assert student.post(f"/api/applications/{claim['id']}/submit", json={}).status_code == 200
    reviewer = login("reviewer1")
    rejection = reviewer.post(f"/api/applications/{claim['id']}/review", json={"action": "REJECT", "reason": "MISSING_VALIDITY"})
    assert rejection.status_code == 200, rejection.text
    response = student.delete(f"/api/applications/{claim['id']}/crops/{crop_id}")
    assert response.status_code == 200, response.text
    detail = student.get(f"/api/applications/{claim['id']}").json()
    assert detail["crops"] == []
    with db.read() as conn:
        historical_crop = db.one(conn, "evidence_crops", crop_id)
        assert historical_crop is not None
        frozen_revision = db.one(conn, "application_revisions", historical_crop["revision_id"])
        assert frozen_revision["content_frozen_at"]


def test_successful_analysis_requires_explicit_apply_and_keeps_manual_override(material_app, monkeypatch):
    _, login = material_app
    student = login("student")
    claim, _, _ = create_claim(student)
    record = upload(student, claim["id"])
    page_id = record["pages"][0]["id"]
    initial_crop = student.post(f"/api/applications/{claim['id']}/crops", json={"page_id": page_id, "evidence_type": "PARTICIPATION", "box": [0, 0, 200, 100]})
    assert initial_crop.status_code == 200
    config = {"enabled": True, "secret_id": "test", "secret_key": "test", "region": "ap-shanghai", "deepseek_key": "test", "deepseek_model": "test-no-network", "general_limit": 10, "seal_limit": 10, "max_analysis_pages": 12}
    monkeypatch.setattr(materials, "configuration", lambda: config)

    def fake_ocr(action, image, config):
        if action == "SealOCR":
            return {"RequestId": "test-no-network-seal", "SealInfos": []}
        return {"RequestId": "test-no-network-text", "TextDetections": [{"DetectedText": "王欣然参加仲英书院周末广场舞（2次），2026年10月24日，学号202401004", "Confidence": 100, "Polygon": [{"X": 15, "Y": 20}, {"X": 190, "Y": 20}, {"X": 190, "Y": 45}, {"X": 15, "Y": 45}]}]}

    monkeypatch.setattr(materials, "tencent_ocr", fake_ocr)
    monkeypatch.setattr(materials, "deepseek_suggestion", lambda context, config: {"title": "仲英书院周末广场舞（2次）", "requested_score": "0.60", "activity_id": None, "payload_json": {}, "explanation": "每次0.3分，两次合并一个条目。"})
    result = student.post(f"/api/applications/{claim['id']}/analyze", json={})
    assert result.status_code == 200, result.text
    assert result.json()["status"] == "SUCCEEDED"
    assert result.json()["analysis"]["id"]
    before_apply = student.get(f"/api/applications/{claim['id']}").json()
    assert before_apply["revision"]["requested_score"] == "0.00"
    applied = student.post(f"/api/applications/{claim['id']}/apply-analysis", json={"analysis_id": result.json()["analysis_id"]})
    assert applied.status_code == 200, applied.text
    assert applied.json()["revision"]["requested_score"] == "0.60"
    assert applied.json()["revision"]["student_modified_score"] is False
    manual = student.put(f"/api/applications/{claim['id']}", json={"title": "学生核对后修订", "requested_score": "0.90", "payload_json": {}})
    assert manual.status_code == 200, manual.text
    assert manual.json()["revision"]["student_modified_score"] is True
    reverted = student.put(f"/api/applications/{claim['id']}", json={"title": "学生核对后修订", "requested_score": "0.60", "payload_json": {}})
    assert reverted.status_code == 200
    assert reverted.json()["revision"]["student_modified_score"] is True


def test_export_has_original_fifteen_columns_all_students_and_no_formulas(material_app):
    _, login = material_app
    monitor = login("monitor")
    bootstrap = monitor.get("/api/bootstrap").json()
    class_id, year_id = bootstrap["classes"][0]["id"], bootstrap["years"][0]["id"]
    response = monitor.get("/api/export", params={"class_id": class_id, "year_id": year_id})
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_column == 15
    assert "E1:I1" in [str(r) for r in sheet.merged_cells.ranges]
    assert "J1:L1" in [str(r) for r in sheet.merged_cells.ranges]
    assert sheet.max_row == 7  # All five demo classmates, including no-claim students.
    assert all(sheet.cell(row, 5).value == 70 and sheet.cell(row, 9).value == 0 and sheet.cell(row, 14).value == 70 for row in range(3, 8))
    assert all(sheet.cell(row, 15).value is None for row in range(3, 8))
    assert login("student").get("/api/export", params={"class_id": class_id, "year_id": year_id}).status_code == 403
    assert materials.excel_text("=HYPERLINK(\"https://example.invalid\")").startswith("'=")


def _text_run(conn, page, texts, api="GeneralAccurateOCR"):
    run = db.add(conn, "recognition_runs", page_id=page["id"], provider="TENCENT", api_name=api, parameters_hash="synthetic-only", status="SUCCEEDED", result_json={}, started_at=db.now(), finished_at=db.now())
    for index, (text, box) in enumerate(texts):
        db.add(conn, "recognition_regions", run_id=run["id"], region_no=index, kind="SEAL" if api == "SealOCR" else "TEXT", text=text, confidence=Decimal("1") if api != "SealOCR" else None, polygon_json=materials._polygon(box), extra_json={})
    return run


def test_v2_one_context_crop_prefers_exact_student_number_and_preserves_source_pages(material_app):
    _, login = material_app
    client = login("student")
    claim, _, _ = create_claim(client)
    pdf = io.BytesIO()
    Image.new("RGB", (400, 800), "white").save(pdf, format="PDF", save_all=True, append_images=[Image.new("RGB", (400, 800), "white")])
    response = client.post(f"/api/applications/{claim['id']}/files", files={"file": ("two-pages.pdf", pdf.getvalue(), "application/pdf")})
    assert response.status_code == 200, response.text
    file_id = response.json()["id"]
    with db.write() as conn:
        user = db.one(conn, "users", username="student")
        context = materials._app_context(conn, claim["id"])
        revision = db.one(conn, "application_revisions", context["application"]["working_revision_id"])
        pages = sorted(db.rows(conn, "file_pages", file_id=file_id), key=lambda page: page["page_index"])
        _text_run(conn, pages[0], [("仲英书院周末广场舞活动", [30, 30, 400, 30]), ("2026年10月24日", [30, 100, 300, 30]), ("特此证明", [30, 180, 200, 30]), (user["display_name"], [30, 700, 150, 30]), (context["class"]["name"], [30, 1000, 180, 30])])
        _text_run(conn, pages[0], [("仲英书院", [500, 1200, 180, 180])], api="SealOCR")
        number_box = [40, 1200, 300, 30]
        _text_run(conn, pages[1], [(context["student"]["student_number"], number_box)])
        suggestions = materials._evidence_suggestions(conn, pages, context["student"], user, context["class"]["name"])
        primary = [item for item in suggestions if item["evidence_type"] == "PARTICIPATION"]
        assert len(primary) == 1
        primary = primary[0]
        assert primary["page_id"] == pages[1]["id"]
        assert primary["page_number"] == 2
        assert primary["file_name"] == "two-pages.pdf"
        assert primary["box"][0] == 0 and primary["box"][2] == pages[1]["width_px"]
        assert primary["box"][3] == round(pages[1]["height_px"] / 4)
        assert primary["highlights"][0]["field"] == "STUDENT_NUMBER"
        assert any(match["page_id"] == pages[0]["id"] and match["field"] == "NAME" for match in primary["other_matches"])
        assert {item["validity_kind"] for item in suggestions if item["evidence_type"] == "VALIDITY"} == {"NAME", "TIME", "SEAL"}
        # The generic '特此证明' row is not a fourth kind of validity evidence.
        assert not any(item["box"][1] == 172 for item in suggestions if item["evidence_type"] == "VALIDITY")
        fingerprint = materials.revision_fingerprint(conn, revision)
        result = {"input_revision_id": revision["id"], "input_fingerprint": fingerprint, "suggested_crops": suggestions}
    applied = materials.apply_ocr_evidence(claim["id"], user, result, expected_fingerprint=fingerprint)
    assert applied["applied"] is True
    detail = client.get(f"/api/applications/{claim['id']}").json()
    crops = [crop for crop in detail["crops"] if crop["evidence_type"] == "PARTICIPATION"]
    assert len(crops) == 1 and crops[0]["page_number"] == 2
    assert crops[0]["other_matches"]
    with db.read() as conn:
        crop = db.one(conn, "evidence_crops", crops[0]["id"])
        box = materials.polygon_box(crop["polygon_json"])
        image = Image.open(materials.asset_path(crop["private_asset_uri"]))
        assert image.size == (pages[1]["width_px"], round(pages[1]["height_px"] / 4))
        assert image.getpixel((number_box[0], number_box[1] - box[1])) == (224, 36, 36)
        source = Image.open(materials.asset_path(pages[1]["image_uri"]))
        assert source.getpixel((number_box[0], number_box[1])) == (255, 255, 255)
        refreshed = materials._build_public_crop(conn, crop, pages[1], [])
        assert refreshed["masks_json"]["evidence"]["other_matches"] == crop["masks_json"]["evidence"]["other_matches"]
    manual = client.post(f"/api/applications/{claim['id']}/crops", json={"page_id": pages[0]["id"], "evidence_type": "PARTICIPATION", "box": [50, 650, 200, 200]})
    assert manual.status_code == 200, manual.text
    updated = client.get(f"/api/applications/{claim['id']}").json()
    assert len([crop for crop in updated["crops"] if crop["evidence_type"] == "PARTICIPATION"]) == 1
    assert any(match["page_id"] == pages[1]["id"] for match in manual.json()["other_matches"])


def test_v2_same_item_duplicate_reuses_file_other_item_conflicts_cross_category_allowed(material_app):
    _, login = material_app
    client = login("student")
    first, year, _ = create_claim(client)
    uploaded = upload(client, first["id"])
    renamed = client.post(f"/api/applications/{first['id']}/files", files={"file": ("renamed.png", png(), "image/png")})
    assert renamed.status_code == 200 and renamed.json()["id"] == uploaded["id"]
    assert renamed.json()["duplicate"] is True
    second, _, _ = create_claim(client)
    duplicate = client.post(f"/api/applications/{second['id']}/files", files={"file": ("another-name.png", png(), "image/png")})
    assert duplicate.status_code == 409, duplicate.text
    assert first["id"] in duplicate.text
    bootstrap = client.get("/api/bootstrap").json()
    ideology = next(category["id"] for category in bootstrap["categories"] if category["code"] == "IDEOLOGY")
    other = client.post("/api/applications", json={"category_id": ideology, "year_id": year})
    assert other.status_code == 200
    repeated_category = client.post(f"/api/applications/{other.json()['id']}/files", files={"file": ("proof.png", png(), "image/png")})
    assert repeated_category.status_code == 200, repeated_category.text
    with db.read() as conn:
        assert len(db.rows(conn, "revision_files", revision_id=client.get(f"/api/applications/{first['id']}").json()["revision"]["id"])) == 1


def test_v2_render_cache_reuse_does_not_consume_second_cloud_budget(material_app, monkeypatch):
    _, login = material_app
    client = login("student")
    first, year, _ = create_claim(client)
    uploaded = upload(client, first["id"])
    bootstrap = client.get("/api/bootstrap").json()
    category = next(category["id"] for category in bootstrap["categories"] if category["code"] == "IDEOLOGY")
    other = client.post("/api/applications", json={"year_id": year, "category_id": category}).json()
    second = upload(client, other["id"])
    calls = []
    monkeypatch.setattr(materials, "tencent_ocr", lambda action, image, config: calls.append(action) or {"RequestId": "synthetic-cache", "TextDetections": [{"DetectedText": "周末广场舞", "Confidence": 100, "Polygon": [{"X": 0, "Y": 0}, {"X": 150, "Y": 0}, {"X": 150, "Y": 20}, {"X": 0, "Y": 20}]}]})
    with db.read() as conn:
        first_page = db.one(conn, "file_pages", uploaded["pages"][0]["id"])
        second_page = db.one(conn, "file_pages", second["pages"][0]["id"])
    config = {"general_limit": 1, "seal_limit": 0}
    materials._ocr_page(first_page, "GeneralAccurateOCR", config)
    cached = materials._ocr_page(second_page, "GeneralAccurateOCR", config)
    assert cached["provider"] == "CACHE" and cached["page_id"] == second_page["id"]
    assert calls == ["GeneralAccurateOCR"]
    with db.read() as conn:
        assert materials._usage(conn, "GeneralAccurateOCR") == 1
        assert db.rows(conn, "recognition_regions", run_id=cached["id"])[0]["text"] == "周末广场舞"


def test_v2_ocr_only_reports_progress_and_staged_items_cannot_start_recognition(material_app, monkeypatch):
    _, login = material_app
    client = login("student")
    claim, _, _ = create_claim(client)
    upload(client, claim["id"])
    config = {"enabled": True, "secret_id": "synthetic", "secret_key": "synthetic", "deepseek_key": "synthetic", "deepseek_model": "synthetic", "general_limit": 10, "seal_limit": 10, "max_analysis_pages": 12}
    monkeypatch.setattr(materials, "configuration", lambda: config)
    calls = []

    def fake_ocr(action, image, config):
        calls.append(action)
        return {"RequestId": "synthetic", "TextDetections": [{"DetectedText": "周末广场舞，2026年10月24日，王欣然", "Confidence": 100, "Polygon": [{"X": 10, "Y": 10}, {"X": 200, "Y": 10}, {"X": 200, "Y": 35}, {"X": 10, "Y": 35}]}]} if action == "GeneralAccurateOCR" else {"RequestId": "synthetic-seal", "SealInfos": []}

    monkeypatch.setattr(materials, "tencent_ocr", fake_ocr)
    monkeypatch.setattr(materials, "deepseek_suggestion", lambda *args, **kwargs: pytest.fail("OCR-only must not call DeepSeek"))
    with db.read() as conn:
        user = db.one(conn, "users", username="student")
    progress = []
    result = materials.run_ocr(claim["id"], user, lambda stage, done, total: progress.append((stage, done, total)))
    assert result["status"] == "SUCCEEDED" and result["input_fingerprint"]
    assert calls == ["GeneralAccurateOCR", "SealOCR"]
    assert progress[0] == ("检查PDF已有文字层", 0, 0)
    assert next(item for item in progress if item[0] == "OCR准备")[1:] == (0, 2)
    assert progress[-1][1:] == (2, 2)
    assert client.get(f"/api/applications/{claim['id']}").json()["revision"]["requested_score"] == "0.00"
    with db.read() as conn:
        assert not db.rows(conn, "ai_analyses", application_id=claim["id"])
    category_id = client.get("/api/bootstrap").json()["categories"][0]["id"]
    group = client.post("/api/submissions", json={"category_id": category_id}).json()
    staged = client.post(f"/api/submissions/{group['id']}/items", json={"staged": True}).json()
    pending = client.post(f"/api/applications/{staged['id']}/analyze", json={})
    assert pending.status_code == 409
    assert calls == ["GeneralAccurateOCR", "SealOCR"]
    saved = client.put(f"/api/applications/{claim['id']}", json={"title": "学生自己修改", "requested_score": "0.90", "payload_json": {}})
    assert saved.status_code == 200
    with pytest.raises(HTTPException) as stale:
        materials.apply_ocr_evidence(claim["id"], user, result, expected_fingerprint=result["input_fingerprint"])
    assert stale.value.status_code == 409


def test_v2_noncurrent_year_is_read_only_for_all_student_material_mutations(material_app):
    _, login = material_app
    student = login("student")
    claim, year_id, _ = create_claim(student)
    uploaded = upload(student, claim["id"])
    admin = login("admin")
    other_year = next(year["id"] for year in admin.get("/api/bootstrap").json()["years"] if year["id"] != year_id)
    switched = admin.put("/api/admin/statistics-year", json={"year_id": other_year})
    assert switched.status_code == 200, switched.text
    assert student.get(uploaded["url"]).status_code == 200
    assert student.post(f"/api/applications/{claim['id']}/files", files={"file": ("new.png", png(color="blue"), "image/png")}).status_code == 409
    assert student.delete(f"/api/applications/{claim['id']}/files/{uploaded['id']}").status_code == 409
    assert student.post(f"/api/applications/{claim['id']}/crops", json={"page_id": uploaded["pages"][0]["id"], "evidence_type": "VALIDITY", "validity_kind": "NAME", "box": [0, 0, 100, 100]}).status_code == 409
    assert student.post(f"/api/applications/{claim['id']}/analyze", json={}).status_code == 409


def test_v2_generic_activity_titles_are_rejected():
    assert not materials.specific_activity_title("活动证明")
    assert not materials.specific_activity_title("2026年两次集体活动证明")
    assert not materials._name_region("特此证明", [])
    assert not materials._name_region("兹有学生参加上述活动", [])
    assert materials.specific_activity_title("仲英书院周末广场舞（2次）")
    assert materials._name_region("活动名称：仲英书院夜跑活动", [])
    assert materials.activity_title_supported("仲英书院周末广场舞（2次）", [{"text": "2026年周末广场舞活动"}], [])
    assert not materials.activity_title_supported("完全虚构的飞天竞赛", [{"text": "王欣然参加周末广场舞"}], [])


def test_ocr_source_fingerprint_ignores_form_edits_but_tracks_attachment_pages(material_app):
    _, login = material_app
    client = login("student")
    claim, _, _ = create_claim(client)
    uploaded = upload(client, claim["id"])
    with db.read() as conn:
        application = db.one(conn, "applications", claim["id"])
        revision = db.one(conn, "application_revisions", application["working_revision_id"])
        source_before = materials.ocr_input_fingerprint(conn, revision)
        full_before = materials.revision_fingerprint(conn, revision)
    saved = client.put(f"/api/applications/{claim['id']}", json={"title": "周末广场舞", "requested_score": "0.60", "payload_json": {"note": "学生改稿"}})
    assert saved.status_code == 200
    cropped = client.post(f"/api/applications/{claim['id']}/crops", json={"page_id": uploaded["pages"][0]["id"], "evidence_type": "PARTICIPATION", "box": [0, 20, 100, 80]})
    assert cropped.status_code == 200
    with db.read() as conn:
        revision = db.one(conn, "application_revisions", revision["id"])
        assert materials.ocr_input_fingerprint(conn, revision) == source_before
        assert materials.revision_fingerprint(conn, revision) != full_before
    added = client.post(f"/api/applications/{claim['id']}/files", files={"file": ("another.png", png(color="blue"), "image/png")})
    assert added.status_code == 200
    with db.write() as conn:
        source_after = materials.ocr_input_fingerprint(conn, revision)
        assert source_after != source_before
        page = db.one(conn, "file_pages", added.json()["pages"][0]["id"])
        db.change(conn, "file_pages", page["id"], render_sha256="changed-source-hash")
        assert materials.ocr_input_fingerprint(conn, revision) != source_after


def test_ai_dependency_reuses_successful_ocr_without_retrying_failed_seal(material_app, monkeypatch):
    _, login = material_app
    client = login("student")
    claim, _, _ = create_claim(client)
    upload(client, claim["id"])
    config = {"enabled": True, "secret_id": "synthetic", "secret_key": "synthetic", "deepseek_key": "synthetic", "deepseek_model": "synthetic", "general_limit": 10, "seal_limit": 10, "max_analysis_pages": 12}
    monkeypatch.setattr(materials, "configuration", lambda: config)
    calls = []

    def fake_ocr(action, image, config):
        calls.append(action)
        if action == "SealOCR":
            raise ProviderError("FailedOperation.ImageNoText", "未识别出印章。")
        return {"RequestId": "synthetic", "TextDetections": [{"DetectedText": "王欣然参加周末广场舞，2026年10月24日", "Confidence": 100, "Polygon": [{"X": 10, "Y": 10}, {"X": 200, "Y": 10}, {"X": 200, "Y": 35}, {"X": 10, "Y": 35}]}]}

    monkeypatch.setattr(materials, "tencent_ocr", fake_ocr)
    with db.read() as conn:
        user = db.one(conn, "users", username="student")
    ocr = materials.run_ocr(claim["id"], user)
    assert ocr["status"] == "SUCCEEDED" and ocr["ocr_fingerprint"]
    assert calls == ["GeneralAccurateOCR", "SealOCR"]
    monkeypatch.setattr(materials, "run_ocr", lambda *args, **kwargs: pytest.fail("AI dependency must not rerun OCR"))
    monkeypatch.setattr(materials, "deepseek_suggestion", lambda *args, **kwargs: {"title": "周末广场舞", "requested_score": "0.30", "activity_id": None, "payload_json": {}})
    result = materials.run_ai(claim["id"], user, ocr_result=ocr)
    assert result["status"] == "SUCCEEDED"
    assert calls == ["GeneralAccurateOCR", "SealOCR"]
    with db.read() as conn:
        assert materials._usage(conn, "SealOCR") == 1
        assert len(db.rows(conn, "recognition_runs", api_name="SealOCR")) == 1
    saved = client.put(f"/api/applications/{claim['id']}", json={"title": "周末广场舞（学生修订）", "requested_score": "0.60", "payload_json": {}})
    assert saved.status_code == 200
    with pytest.raises(HTTPException) as stale:
        materials.run_ai(claim["id"], user, ocr_result=ocr)
    assert stale.value.status_code == 409
    assert calls == ["GeneralAccurateOCR", "SealOCR"]


def test_ocr_evidence_source_guard_and_queue_advance_are_one_transaction(material_app):
    _, login = material_app
    client = login("student")
    claim, _, _ = create_claim(client)
    uploaded = upload(client, claim["id"])
    with db.write() as conn:
        user = db.one(conn, "users", username="student")
        context = materials._app_context(conn, claim["id"])
        revision = db.one(conn, "application_revisions", context["application"]["working_revision_id"])
        page = db.one(conn, "file_pages", uploaded["pages"][0]["id"])
        run = _text_run(conn, page, [("周末广场舞", [10, 10, 150, 20]), (user["display_name"], [10, 60, 100, 20])])
        original_full = materials.revision_fingerprint(conn, revision)
        source_fp = materials.ocr_input_fingerprint(conn, revision)
        result = {"status": "SUCCEEDED", "application_id": claim["id"], "input_revision_id": revision["id"], "input_fingerprint": original_full, "ocr_fingerprint": source_fp, "page_ids": [page["id"]], "run_ids": [run["id"]], "suggested_crops": []}
    # This edit must not cancel OCR: no attachment has changed.
    assert client.put(f"/api/applications/{claim['id']}", json={"title": "学生的具体活动名", "requested_score": "0.90", "payload_json": {}}).status_code == 200
    with db.write() as conn:
        revision = db.one(conn, "application_revisions", revision["id"])
        before = materials.revision_fingerprint(conn, revision)
        assert before != original_full
        queued = db.add(conn, "background_jobs", application_id=claim["id"], owner_id=user["id"], kind="AI", state="QUEUED", stage="等待OCR", input_revision_id=revision["id"], input_fingerprint=before, result_json={})

    def advance_and_fail(conn, before_fp, after_fp):
        assert before_fp == before and after_fp != before_fp
        db.change(conn, "background_jobs", queued["id"], input_fingerprint=after_fp)
        raise RuntimeError("synthetic failure after queue update")

    with pytest.raises(RuntimeError, match="synthetic failure"):
        materials.apply_ocr_evidence(claim["id"], user, result, expected_fingerprint=original_full, expected_ocr_fingerprint=source_fp, on_applied=advance_and_fail)
    with db.read() as conn:
        assert not db.rows(conn, "evidence_crops", revision_id=revision["id"])
        assert db.one(conn, "background_jobs", queued["id"])["input_fingerprint"] == before
    calls = []

    def advance(conn, before_fp, after_fp):
        calls.append((before_fp, after_fp))
        db.change(conn, "background_jobs", queued["id"], input_fingerprint=after_fp)

    applied = materials.apply_ocr_evidence(claim["id"], user, result, expected_fingerprint=original_full, expected_ocr_fingerprint=source_fp, on_applied=advance)
    assert applied["applied"] is True
    with db.read() as conn:
        assert db.one(conn, "background_jobs", queued["id"])["input_fingerprint"] == applied["fingerprint"]
        unchanged_content = db.one(conn, "application_revisions", revision["id"])
        assert unchanged_content["title"] == "学生的具体活动名"
        assert unchanged_content["requested_score"] == Decimal("0.90")
        assert materials.ocr_input_fingerprint(conn, unchanged_content) == source_fp
    assert calls == [(before, applied["fingerprint"])]
    # Actual source changes still invalidate the completed OCR snapshot.
    assert client.post(f"/api/applications/{claim['id']}/files", files={"file": ("new.png", png(color="blue"), "image/png")}).status_code == 200
    with pytest.raises(HTTPException) as stale:
        materials.apply_ocr_evidence(claim["id"], user, result, expected_ocr_fingerprint=source_fp)
    assert stale.value.status_code == 409


def test_last_seal_is_selected_per_file_by_page_y_x_and_old_results_are_filtered(material_app):
    _, login = material_app
    client = login("student")
    claim, _, _ = create_claim(client)
    pdf = io.BytesIO()
    Image.new("RGB", (500, 800), "white").save(pdf, format="PDF", save_all=True, append_images=[Image.new("RGB", (500, 800), "white")])
    response = client.post(f"/api/applications/{claim['id']}/files", files={"file": ("multi-seals.pdf", pdf.getvalue(), "application/pdf")})
    assert response.status_code == 200, response.text
    second = client.post(f"/api/applications/{claim['id']}/files", files={"file": ("supplement.png", png(color="blue"), "image/png")})
    assert second.status_code == 200
    with db.write() as conn:
        user = db.one(conn, "users", username="student")
        context = materials._app_context(conn, claim["id"])
        revision = db.one(conn, "application_revisions", context["application"]["working_revision_id"])
        pdf_pages = sorted(db.rows(conn, "file_pages", file_id=response.json()["id"]), key=lambda page: page["page_index"])
        supplement = db.one(conn, "file_pages", second.json()["pages"][0]["id"])
        old_suggestions = []
        expected_regions = []

        def seal_run(page, locations, expected_text):
            data = {"RequestId": "synthetic-seals", "SealInfos": [{"SealBody": text, "Location": {"X": x, "Y": y, "Width": 80, "Height": 60}} for text, x, y in locations]}
            run = db.add(conn, "recognition_runs", page_id=page["id"], provider="TENCENT", api_name="SealOCR", parameters_hash="synthetic-seals", status="SUCCEEDED", result_json=data, started_at=db.now(), finished_at=db.now())
            materials._save_regions(conn, run, data)
            for region in db.rows(conn, "recognition_regions", run_id=run["id"]):
                x, y, w, h = materials.polygon_box(region["polygon_json"])
                old_suggestions.append({"page_id": page["id"], "evidence_type": "VALIDITY", "validity_kind": "SEAL", "region_id": region["id"], "box": [x - 5, y - 5, w + 10, h + 10]})
                if region["text"] == expected_text:
                    expected_regions.append(region["id"])

        # Last element in provider order is deliberately not the last stamp on the page.
        seal_run(pdf_pages[0], [("first-page-bottom", 850, 1450), ("first-page-top", 20, 50)], None)
        seal_run(pdf_pages[1], [("last-pdf", 700, 600), ("same-y-left", 500, 600), ("last-array-but-earlier", 800, 100)], "last-pdf")
        seal_run(supplement, [("last-image", 170, 120), ("same-y-left", 20, 120), ("last-array-but-earlier", 200, 30)], "last-image")
        pages = [supplement, pdf_pages[1], pdf_pages[0]]  # Input page order is deliberately mixed too.
        suggestions = materials._evidence_suggestions(conn, pages, context["student"], user, context["class"]["name"])
        seals = [item for item in suggestions if item.get("validity_kind") == "SEAL"]
        assert len(seals) == 2
        assert {item["region_id"] for item in seals} == set(expected_regions)
        full_fp = materials.revision_fingerprint(conn, revision)
        result = {"status": "SUCCEEDED", "input_revision_id": revision["id"], "input_fingerprint": full_fp, "ocr_fingerprint": materials.ocr_input_fingerprint(conn, revision), "page_ids": [page["id"] for page in pages], "suggested_crops": list(reversed(old_suggestions))}
    # Even a persisted job carrying every stamp is re-derived from existing OCR, without a cloud call.
    materials.apply_ocr_evidence(claim["id"], user, result, expected_fingerprint=full_fp)
    with db.write() as conn:
        current = db.rows(conn, "evidence_crops", revision_id=revision["id"])
        assert len(current) == 2 and {crop["recognition_region_id"] for crop in current} == set(expected_regions)
        # The compatibility path without page_ids also filters old stored AI suggestions.
        earlier = next(item for item in old_suggestions if item["region_id"] not in expected_regions)
        page = db.one(conn, "file_pages", earlier["page_id"])
        materials.create_crop(conn, revision, page, "VALIDITY", earlier["box"], source="AI", region_id=earlier["region_id"], metadata={"validity_kind": "SEAL"})
        assert len(db.rows(conn, "evidence_crops", revision_id=revision["id"])) == 3
        materials._apply_evidence(conn, revision, list(reversed(old_suggestions)))
        final = db.rows(conn, "evidence_crops", revision_id=revision["id"])
        assert len(final) == 2 and {crop["recognition_region_id"] for crop in final} == set(expected_regions)
        frozen_ids = {crop["id"] for crop in final}
    assert client.put(f"/api/applications/{claim['id']}", json={"title": "周末广场舞", "requested_score": "0.30", "payload_json": {}}).status_code == 200
    assert client.post(f"/api/applications/{claim['id']}/submit", json={}).status_code == 200
    assert login("reviewer1").post(f"/api/applications/{claim['id']}/review", json={"action": "APPROVE"}).status_code == 200
    with pytest.raises(HTTPException) as locked:
        materials.apply_ocr_evidence(claim["id"], user, result, expected_ocr_fingerprint=result["ocr_fingerprint"])
    assert locked.value.status_code == 409
    with db.read() as conn:
        assert {crop["id"] for crop in db.rows(conn, "evidence_crops", revision_id=revision["id"])} == frozen_ids

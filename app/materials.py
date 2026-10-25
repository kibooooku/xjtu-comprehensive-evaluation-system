"""Protected evidence storage, OCR orchestration and the original 15-column export."""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import threading
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pypdfium2 as pdfium
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from PIL import Image, ImageDraw, ImageOps, UnidentifiedImageError
from sqlalchemy import and_, delete, func, select

from . import db
from .db import add, change, money, now, one, rows, tables, uid
from .integrations import ProviderError, configuration, deepseek_suggestion, tencent_ocr
from .limits import MAX_ATTACHMENT_BYTES, MAX_ATTACHMENT_PAGES, MAX_AUTOMATIC_OCR_PAGES, OCR_POLICY_VERSION
from .security import current_user, require

router = APIRouter(prefix="/api")
MAX_UPLOAD_BYTES = MAX_ATTACHMENT_BYTES
MAX_PAGES = MAX_ATTACHMENT_PAGES
MAX_PAGE_PIXELS = 36_000_000
REDACTION_VERSION = "student-number-v1"
PDF_LOCK = threading.Lock()
ANALYSIS_LOCK = threading.Lock()
OCR_CALL_LOCK = threading.Lock()
Image.MAX_IMAGE_PIXELS = MAX_PAGE_PIXELS
NATIVE_TEXT_API = "NativePDFText"


def _page_limit(config=None):
    config = configuration() if config is None else config
    try:
        return max(1, min(MAX_AUTOMATIC_OCR_PAGES, int(config.get("max_analysis_pages", MAX_AUTOMATIC_OCR_PAGES))))
    except (TypeError, ValueError):
        return MAX_AUTOMATIC_OCR_PAGES


def asset_path(relative: str) -> Path:
    root = Path(db.storage_root).resolve()
    path = (root / relative).resolve()
    if not path.is_relative_to(root):
        raise HTTPException(404, "文件不存在。")
    return path


def _write_asset(relative: str, data: bytes) -> Path:
    path = asset_path(relative)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def _png(image: Image.Image) -> bytes:
    stream = io.BytesIO()
    image.save(stream, format="PNG", optimize=True)
    return stream.getvalue()


def render_upload(content: bytes, filename: str) -> tuple[str, list[bytes]]:
    """Decode actual file contents, never trust filename or client MIME type."""
    suffix = Path(filename).suffix.lower()
    if suffix not in (".pdf", ".jpg", ".jpeg", ".png"):
        raise HTTPException(400, "仅支持 PDF、JPG、JPEG、PNG 文件。")
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(400, "每个文件需大于0字节且不超过200MB。")
    pages = []
    if suffix == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise HTTPException(400, "文件内容不是有效PDF。")
        try:
            with PDF_LOCK:
                document = pdfium.PdfDocument(content)
                try:
                    if not 1 <= len(document) <= MAX_PAGES:
                        raise HTTPException(400, f"PDF需包含1至{MAX_PAGES}页。")
                    for index in range(len(document)):
                        page = document[index]
                        width, height = page.get_size()
                        if width <= 0 or height <= 0:
                            raise HTTPException(400, "PDF页面尺寸无效。")
                        scale = min(2.0, 3600 / max(width, height))
                        bitmap = page.render(scale=scale)
                        image = bitmap.to_pil().convert("RGB")
                        pages.append(_png(image))
                        bitmap.close()
                        page.close()
                finally:
                    document.close()
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(400, "PDF无法解析，可能已损坏或有密码保护。") from None
        return "application/pdf", pages
    try:
        with Image.open(io.BytesIO(content)) as image:
            if image.format not in ("JPEG", "PNG"):
                raise HTTPException(400, "图片内容与支持的格式不符。")
            if (suffix == ".png") != (image.format == "PNG"):
                raise HTTPException(400, "图片扩展名与实际内容不符。")
            if image.width * image.height > MAX_PAGE_PIXELS:
                raise HTTPException(400, "图片分辨率过大，请压缩后上传。")
            mime = "image/png" if image.format == "PNG" else "image/jpeg"
            normalized = ImageOps.exif_transpose(image).convert("RGB")
            if normalized.width < 10 or normalized.height < 10:
                raise HTTPException(400, "图片尺寸过小。")
            pages.append(_png(normalized))
    except HTTPException:
        raise
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
        raise HTTPException(400, "图片无法解码或尺寸过大。") from None
    return mime, pages


def file_view(conn, record: dict) -> dict:
    pages = sorted(rows(conn, "file_pages", file_id=record["id"]), key=lambda page: page["page_index"])
    native_pages = [page["page_index"] for page in pages if _native_text_cached(conn, page)]
    max_pages = _page_limit()
    long_file = len(pages) > max_pages
    native_ids = {page["id"] for page in pages if page["page_index"] in native_pages}
    text_ids = set(native_ids)
    cloud_text_ids = set()
    checked_seals, checked_text_layers, seal_found = set(), set(), False
    boundary_ids = {pages[0]["id"], pages[-1]["id"]} if pages else set()
    for page in pages:
        latest = {}
        for run in rows(conn, "recognition_runs", page_id=page["id"]):
            if run["status"] != "SUCCEEDED":
                continue
            if run["api_name"] not in latest or (run.get("finished_at") or "") >= (latest[run["api_name"]].get("finished_at") or ""):
                latest[run["api_name"]] = run
        for run in latest.values():
            if run["api_name"] == NATIVE_TEXT_API and "usable" in (run.get("result_json") or {}):
                checked_text_layers.add(page["id"])
            if run["api_name"] == "GeneralAccurateOCR" and (not long_file or page["id"] in boundary_ids) and any((region.get("text") or "").strip() for region in rows(conn, "recognition_regions", run_id=run["id"])):
                text_ids.add(page["id"])
                cloud_text_ids.add(page["id"])
            if run["api_name"] == "SealOCR" and page["id"] in boundary_ids:
                checked_seals.add(page["id"])
                seal_found = seal_found or any(region["kind"] == "SEAL" for region in rows(conn, "recognition_regions", run_id=run["id"]))
    seal_status = "FOUND" if seal_found else ("NOT_FOUND" if boundary_ids and boundary_ids <= checked_seals else ("LIMITED" if checked_seals else "NOT_RUN"))
    text_layer_check_complete = record["mime_type"] != "application/pdf" or len(checked_text_layers) == len(pages)
    return {
        "id": record["id"], "original_name": record["original_name"],
        "mime_type": record["mime_type"], "byte_size": record["byte_size"],
        "url": f"/api/files/{record['id']}/original",
        "page_count": len(pages), "max_pages": max_pages, "cloud_ocr_eligible": True,
        "full_page_cloud_ocr_eligible": not long_file, "cloud_ocr_scope": "ENDPOINTS" if long_file else "FULL",
        "ocr_page_numbers": [page["page_index"] for page in pages if page["id"] in cloud_text_ids], "text_covered_page_count": len(text_ids),
        "participation_mode": ("AUTO" if not long_file or native_pages else "MANUAL") if text_layer_check_complete else None,
        "text_layer_used": bool(native_pages), "native_pages": native_pages,
        "text_layer_check_complete": text_layer_check_complete, "text_layer_checked_page_count": len(checked_text_layers),
        "native_text_status": "PENDING" if not text_layer_check_complete else ("USED" if pages and len(native_pages) == len(pages) else ("PARTIAL" if native_pages else "NONE")),
        "automatic_ocr_status": "CHECKING_TEXT_LAYER" if not text_layer_check_complete else ("COMPLETED" if pages and len(text_ids) == len(pages) else ("PARTIAL" if text_ids else "ELIGIBLE")),
        "ocr_skip_reason": f"超过{max_pages}页：无文字层时仅云OCR首页和末页，中间扫描页不识别；请手动寻找并框选本人参与信息。" if long_file and text_layer_check_complete and any(page["id"] not in native_ids and page["id"] not in boundary_ids for page in pages) else None,
        "seal_pages": sorted({pages[0]["page_index"], pages[-1]["page_index"]}) if pages else [],
        "manual_page_numbers": [page["page_index"] for page in pages if page["id"] not in native_ids and page["id"] not in boundary_ids] if long_file and text_layer_check_complete else [],
        "seal_status": seal_status, "seal_found": seal_found,
        "seal_message": "仅检查文件首页和末页，未定位到印章；请自行寻找并手动框选。" if seal_status in ("NOT_FOUND", "LIMITED") else ("已定位首页/末页中的最后一枚印章，请核对原件。" if seal_found else "将仅检查文件首页和末页。"),
        "pages": [{"id": page["id"], "page_index": page["page_index"], "width_px": page["width_px"], "height_px": page["height_px"], "url": f"/api/pages/{page['id']}/image"} for page in pages],
    }


def polygon_box(polygon) -> list[int]:
    if isinstance(polygon, dict):
        return list(polygon.get("box", [0, 0, 0, 0]))
    if not polygon:
        return [0, 0, 0, 0]
    points = [(p.get("X", p.get("x", 0)), p.get("Y", p.get("y", 0))) if isinstance(p, dict) else tuple(p) for p in polygon]
    x1, y1 = min(p[0] for p in points), min(p[1] for p in points)
    x2, y2 = max(p[0] for p in points), max(p[1] for p in points)
    return [int(x1), int(y1), int(x2 - x1), int(y2 - y1)]


def crop_view(record: dict, public: bool = False, conn=None) -> dict:
    ready = record.get("public_asset_status") == "READY"
    result = {"id": record["id"], "page_id": record["page_id"], "evidence_type": record["evidence_type"], "public_asset_status": record.get("public_asset_status"), "public_url": f"/api/crops/{record['id']}/public" if ready else None}
    if not public:
        metadata = record.get("masks_json") or {}
        evidence = metadata.get("evidence") or {}
        if not evidence.get("file_name") or not evidence.get("page_number"):
            # Frozen V1 evidence has no presentation metadata; derive its real
            # source instead of incorrectly labelling every legacy crop page 1.
            if conn is not None:
                original_page = one(conn, "file_pages", record["page_id"])
                original_file = one(conn, "file_objects", original_page["file_id"]) if original_page else None
            else:
                with db.read() as source_conn:
                    original_page = one(source_conn, "file_pages", record["page_id"])
                    original_file = one(source_conn, "file_objects", original_page["file_id"]) if original_page else None
            evidence = {**evidence, "file_name": evidence.get("file_name") or (original_file or {}).get("original_name"),
                        "page_number": evidence.get("page_number") or (original_page or {}).get("page_index")}
        result.update({"source": record["source"], "box": polygon_box(record["polygon_json"]), "url": f"/api/crops/{record['id']}/image", "masks": metadata.get("manual", []), "file_name": evidence.get("file_name"), "page_number": evidence.get("page_number"), "highlights": evidence.get("highlights", []), "other_matches": evidence.get("other_matches", []), "validity_kind": evidence.get("validity_kind"), "layout_version": evidence.get("layout_version")})
        if conn is not None:
            accepted, reason = automatic_crop_policy(conn, record)
        else:
            with db.read() as policy_conn:
                accepted, reason = automatic_crop_policy(policy_conn, record)
        result.update(automatic_policy_allowed=accepted, policy_exclusion_reason=reason)
        result["origin"] = evidence.get("origin")
    return result


def _file_policy(conn, file_id, max_pages=None):
    pages = sorted(rows(conn, "file_pages", file_id=file_id), key=lambda page: page["page_index"])
    return pages, bool(pages and len(pages) <= (max_pages if max_pages is not None else _page_limit())), {pages[0]["id"], pages[-1]["id"]} if pages else set()


def _native_text_cached(conn, page):
    candidates = [run for run in rows(conn, "recognition_runs", page_id=page["id"], api_name=NATIVE_TEXT_API) if run["status"] == "SUCCEEDED" and (run.get("result_json") or {}).get("usable")]
    return sorted(candidates, key=lambda run: run.get("finished_at") or "")[-1] if candidates else None


def _is_seal_crop(conn, crop):
    metadata = (crop.get("masks_json") or {}).get("evidence") or {}
    region = one(conn, "recognition_regions", crop["recognition_region_id"]) if crop.get("recognition_region_id") else None
    return metadata.get("validity_kind") == "SEAL" or bool(region and region["kind"] == "SEAL")


def automatic_crop_policy(conn, crop, revision=None):
    """Hide obsolete machine suggestions only; never rewrite manual/frozen evidence."""
    if crop.get("source") != "AI":
        return True, None
    revision = revision or one(conn, "application_revisions", crop.get("revision_id"))
    if revision and revision.get("content_frozen_at"):
        return True, None
    page = one(conn, "file_pages", crop.get("page_id"))
    if not page:
        return False, "来源页不存在。"
    max_pages = _page_limit()
    _, cloud_allowed, boundaries = _file_policy(conn, page["file_id"], max_pages)
    if _is_seal_crop(conn, crop):
        return (True, None) if page["id"] in boundaries else (False, "自动查章只限文件首页和末页，中间页旧候选不再展示。")
    if not cloud_allowed:
        region = one(conn, "recognition_regions", crop["recognition_region_id"]) if crop.get("recognition_region_id") else None
        run = one(conn, "recognition_runs", region["run_id"]) if region else None
        if run and run["api_name"] == NATIVE_TEXT_API:
            return True, None
        if crop.get("evidence_type") == "PARTICIPATION":
            return False, f"超过{max_pages}页扫描文件的本人参与信息须手动选择，首末页云OCR仅用于有效性证据。"
        if not run or run["api_name"] != "GeneralAccurateOCR" or page["id"] not in boundaries:
            return False, f"超过{max_pages}页文件的中间扫描页不调用云OCR，旧中间页候选不再自动展示。"
    return True, None


def _app_context(conn, application_id: str) -> dict:
    from .workflow import app_context
    return app_context(conn, application_id)


def _editable(conn, application_id: str, user: dict) -> tuple[dict, dict]:
    from .workflow import editable_revision
    context = _app_context(conn, application_id)
    from .migrations import assert_current_year
    assert_current_year(conn, context["year"]["id"])
    context["previous_revision_id"] = context["application"].get("working_revision_id")
    require(conn, user, "application.write", class_id=context["class"]["id"], owner_id=context["user"]["id"])
    revision = editable_revision(conn, context["application"], user)
    return context, revision


def _can_read_file(conn, file_id: str, user: dict) -> dict:
    record = one(conn, "file_objects", file_id)
    if not record:
        raise HTTPException(404, "文件不存在。")
    if user["base_identity"] == "ADMIN":
        raise HTTPException(403, "管理员不能查看业务证明材料。")
    links = rows(conn, "revision_files", file_id=file_id)
    for link in links:
        revision = one(conn, "application_revisions", link["revision_id"])
        if not revision:
            continue
        from .workflow import app_context
        context = app_context(conn, revision["application_id"], include_archived=True)
        try:
            require(conn, user, "materials.read", class_id=context["class"]["id"], owner_id=context["user"]["id"])
            return record
        except HTTPException as error:
            if error.status_code != 403:
                raise
    raise HTTPException(403, "你没有查看此材料的权限。")


def _private_response(relative: str, media_type: str = "image/png", filename: str | None = None):
    path = asset_path(relative)
    if not path.is_file():
        raise HTTPException(404, "文件不存在。")
    return FileResponse(path, media_type=media_type, filename=filename, headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff", "Referrer-Policy": "no-referrer"})


@router.post("/applications/{application_id}/files")
def upload_file(application_id: str, file: UploadFile = File(...), user: dict = Depends(current_user)):
    content = file.file.read(MAX_UPLOAD_BYTES + 1)
    filename = Path((file.filename or "证明材料").replace("\\", "/")).name[:240]
    mime, rendered = render_upload(content, filename)
    with db.write() as conn:
        context, revision = _editable(conn, application_id, user)
        attached = rows(conn, "revision_files", revision_id=revision["id"])
        digest = hashlib.sha256(content).hexdigest()
        for link in attached:
            existing = one(conn, "file_objects", link["file_id"])
            if existing["sha256"] == digest:
                return {**file_view(conn, existing), "duplicate": True, "message": "此小项已包含相同内容的附件，已复用原文件。"}
        from .submissions import assert_unique_files
        assert_unique_files(conn, application_id, [digest])
        if len(attached) >= 12:
            raise HTTPException(400, "每条申请最多12份附件。")
        file_id = uid()
        original_uri = f"originals/{file_id}{Path(filename).suffix.lower()}"
        _write_asset(original_uri, content)
        record = add(conn, "file_objects", id=file_id, uploaded_by=user["id"], original_name=filename, mime_type=mime, byte_size=len(content), sha256=digest, storage_uri=original_uri)
        for index, page_bytes in enumerate(rendered):
            page_id = uid()
            uri = f"pages/{page_id}.png"
            _write_asset(uri, page_bytes)
            with Image.open(io.BytesIO(page_bytes)) as image:
                width, height = image.size
            add(conn, "file_pages", id=page_id, file_id=file_id, page_index=index + 1, render_version="pdfium-pillow-v1", width_px=width, height_px=height, rotation_degrees=0, image_uri=uri, render_sha256=hashlib.sha256(page_bytes).hexdigest())
        add(conn, "revision_files", revision_id=revision["id"], file_id=file_id, sort_order=len(attached), purpose="SUPPORT", added_by=user["id"], added_at=now())
        _touch(conn, context["application"])
        return file_view(conn, record)


def _touch(conn, application: dict):
    current = one(conn, "applications", application["id"])
    change(conn, "applications", application["id"], lock_version=(current.get("lock_version") or 0) + 1)


@router.delete("/applications/{application_id}/files/{file_id}")
def remove_file(application_id: str, file_id: str, user: dict = Depends(current_user)):
    with db.write() as conn:
        context, revision = _editable(conn, application_id, user)
        link = one(conn, "revision_files", revision_id=revision["id"], file_id=file_id)
        if not link:
            raise HTTPException(404, "附件不存在。")
        page_ids = [p["id"] for p in rows(conn, "file_pages", file_id=file_id)]
        if page_ids:
            conn.execute(delete(tables["evidence_crops"]).where(and_(tables["evidence_crops"].c.revision_id == revision["id"], tables["evidence_crops"].c.page_id.in_(page_ids))))
        conn.execute(delete(tables["revision_files"]).where(tables["revision_files"].c.id == link["id"]))
        _touch(conn, context["application"])
    return {"ok": True}


@router.get("/files/{file_id}/original")
def get_original(file_id: str, user: dict = Depends(current_user)):
    with db.read() as conn:
        record = _can_read_file(conn, file_id, user)
    return _private_response(record["storage_uri"], record["mime_type"], record["original_name"])


@router.get("/pages/{page_id}/image")
def get_page(page_id: str, user: dict = Depends(current_user)):
    with db.read() as conn:
        page = one(conn, "file_pages", page_id)
        if not page:
            raise HTTPException(404, "页面不存在。")
        _can_read_file(conn, page["file_id"], user)
    return _private_response(page["image_uri"])


def checked_box(value, width: int, height: int) -> list[int]:
    if not isinstance(value, (list, tuple)) or len(value) != 4:
        raise HTTPException(400, "截图范围应为[x,y,width,height]。")
    try:
        x, y, w, h = [float(n) for n in value]
    except (ValueError, TypeError):
        raise HTTPException(400, "截图坐标无效。") from None
    if not all(math.isfinite(n) for n in (x, y, w, h)) or w < 2 or h < 2:
        raise HTTPException(400, "截图范围无效或过小。")
    if x < 0 or y < 0 or x + w > width + 1 or y + h > height + 1:
        raise HTTPException(400, "截图范围超出原图。")
    x1, y1, x2, y2 = max(0, int(x)), max(0, int(y)), min(width, math.ceil(x + w)), min(height, math.ceil(y + h))
    return [x1, y1, x2 - x1, y2 - y1]


def _polygon(box: list[int]) -> list[list[int]]:
    x, y, w, h = box
    return [[x, y], [x + w, y], [x + w, y + h], [x, y + h]]


def _intersection(a: list, b: list) -> list | None:
    x, y = max(a[0], b[0]), max(a[1], b[1])
    right, bottom = min(a[0] + a[2], b[0] + b[2]), min(a[1] + a[3], b[1] + b[3])
    return [x, y, right - x, bottom - y] if right > x and bottom > y else None


def _build_public_crop(conn, crop: dict, page: dict, manual_masks: list) -> dict:
    """Fail closed without a successful text run; never substitute a private crop."""
    metadata = dict(crop.get("masks_json") or {})
    metadata.pop("reason", None)
    runs = [r for r in rows(conn, "recognition_runs", page_id=page["id"]) if r["api_name"] in ("GeneralAccurateOCR", NATIVE_TEXT_API) and r["status"] == "SUCCEEDED" and (r["api_name"] != NATIVE_TEXT_API or (r.get("result_json") or {}).get("usable"))]
    if not runs:
        return {"public_asset_status": "BLOCKED", "public_asset_uri": None, "redaction_version": None, "masks_json": {**metadata, "manual": manual_masks, "automatic": [], "reason": "需要成功的文字OCR定位学号后才能生成公示截图"}}
    run = sorted(runs, key=lambda item: (item["api_name"] == NATIVE_TEXT_API, item.get("finished_at") or ""))[-1]
    regions = rows(conn, "recognition_regions", run_id=run["id"])
    if not any((region.get("text") or "").strip() for region in regions):
        return {"public_asset_status": "BLOCKED", "public_asset_uri": None, "redaction_version": None, "masks_json": {**metadata, "manual": manual_masks, "automatic": [], "reason": "文字OCR未检测到可靠文本，未生成公示截图"}}
    if run["api_name"] == NATIVE_TEXT_API:
        source_numbers = re.findall(r"[0-9０-９]{6,}", _compact((run.get("result_json") or {}).get("text")))
        positioned = _compact("".join(region.get("text") or "" for region in regions))
        if any(number not in positioned for number in source_numbers):
            return {"public_asset_status": "BLOCKED", "public_asset_uri": None, "redaction_version": None, "masks_json": {**metadata, "manual": manual_masks, "automatic": [], "reason": "原生文字中的个人编号未全部取得可靠坐标，未生成公示截图"}}
    automatic = []
    for region in regions:
        text = re.sub(r"\s+", "", region.get("text") or "")
        # Mask every long numeric identifier, including OCR-separated digits and the whole 学号 row.
        if re.search(r"[0-9０-９]{6,}", text) or "学号" in text or (region.get("confidence") is not None and Decimal(str(region["confidence"])) < Decimal("0.85")):
            box = polygon_box(region["polygon_json"])
            if box[2] <= 0 or box[3] <= 0:
                return {"public_asset_status": "BLOCKED", "public_asset_uri": None, "redaction_version": None, "masks_json": {**metadata, "manual": manual_masks, "automatic": [], "reason": "含个人编号的文字缺少可靠坐标，未生成公示截图"}}
            box = [max(0, box[0] - 5), max(0, box[1] - 4), box[2] + 10, box[3] + 8]
            automatic.append(box)
    box = polygon_box(crop["polygon_json"])
    with Image.open(asset_path(crop["private_asset_uri"])) as original:
        image = original.convert("RGB")
    draw = ImageDraw.Draw(image)
    for mask in manual_masks + automatic:
        overlap = _intersection(box, mask)
        if overlap:
            x, y, w, h = overlap
            draw.rectangle((x - box[0], y - box[1], x - box[0] + w, y - box[1] + h), fill="#253244")
    public_uri = f"public-crops/{crop['id']}-{REDACTION_VERSION}.png"
    _write_asset(public_uri, _png(image))
    return {"public_asset_status": "READY", "public_asset_uri": public_uri, "redaction_version": REDACTION_VERSION, "masks_json": {**metadata, "manual": manual_masks, "automatic": automatic, "ocr_run_id": run["id"], "page_hash": page["render_sha256"]}}


def quarter_page_box(page: dict, y: float) -> list[int]:
    height = max(2, int(round(page["height_px"] / 4)))
    return [0, max(0, min(int(round(y)), page["height_px"] - height)), page["width_px"], height]


def create_crop(conn, revision: dict, page: dict, evidence_type: str, box: list, manual_masks: list | None = None, source: str = "STUDENT", region_id: str | None = None, metadata: dict | None = None) -> dict:
    box = checked_box(box, page["width_px"], page["height_px"])
    metadata = dict(metadata or {})
    if revision.get("content_frozen_at"):
        raise HTTPException(409, "已提交的冻结证据不能改写。")
    if evidence_type == "PARTICIPATION":
        box = quarter_page_box(page, box[1])
        metadata["layout_version"] = "context-quarter-v2"
        # Keep one primary identity context; frozen historical revisions are never touched.
        conn.execute(delete(tables["evidence_crops"]).where(and_(tables["evidence_crops"].c.revision_id == revision["id"], tables["evidence_crops"].c.evidence_type == "PARTICIPATION")))
    elif metadata.get("validity_kind") not in ("NAME", "TIME", "SEAL"):
        raise HTTPException(400, "有效性证据请选择活动名称、活动时间或印章。")
    file = one(conn, "file_objects", page["file_id"])
    metadata.update(file_name=file["original_name"], page_number=page["page_index"])
    highlights = []
    for highlight in metadata.get("highlights", []):
        if highlight.get("field") not in ("STUDENT_NUMBER", "NAME", "CLASS"):
            continue
        highlight_box = checked_box(highlight.get("box"), page["width_px"], page["height_px"])
        if _intersection(box, highlight_box):
            highlights.append({"box": highlight_box, "field": highlight["field"]})
    metadata["highlights"] = highlights
    manual_masks = manual_masks or []
    if len(manual_masks) > 100:
        raise HTTPException(400, "手动遮挡区域过多。")
    manual_masks = [checked_box(mask, page["width_px"], page["height_px"]) for mask in manual_masks]
    with Image.open(asset_path(page["image_uri"])) as image:
        x, y, w, h = box
        cropped = image.convert("RGB").crop((x, y, x + w, y + h))
        draw = ImageDraw.Draw(cropped)
        for highlight in highlights:
            overlap = _intersection(box, highlight["box"])
            if overlap:
                hx, hy, hw, hh = overlap
                draw.rectangle((hx - x, hy - y, hx - x + hw - 1, hy - y + hh - 1), outline="#e02424", width=max(2, min(6, round(page["width_px"] / 500))))
        content = _png(cropped)
    crop_id = uid()
    uri = f"private-crops/{crop_id}.png"
    _write_asset(uri, content)
    existing = rows(conn, "evidence_crops", revision_id=revision["id"])
    crop = add(conn, "evidence_crops", id=crop_id, revision_id=revision["id"], page_id=page["id"], evidence_type=evidence_type, source=source, recognition_region_id=region_id, polygon_json=_polygon(box), private_asset_uri=uri, source_hash=hashlib.sha256(content).hexdigest(), redaction_version=None, masks_json={"manual": manual_masks, "evidence": metadata}, public_asset_uri=None, public_asset_status="BLOCKED", sort_order=len(existing))
    return change(conn, "evidence_crops", crop_id, **_build_public_crop(conn, crop, page, manual_masks))


@router.post("/applications/{application_id}/crops")
def add_crop(application_id: str, body: dict, user: dict = Depends(current_user)):
    with db.write() as conn:
        context, revision = _editable(conn, application_id, user)
        existing_crops = rows(conn, "evidence_crops", revision_id=revision["id"])
        replacing_primary = body.get("evidence_type") == "PARTICIPATION" and any(crop["evidence_type"] == "PARTICIPATION" for crop in existing_crops)
        if len(existing_crops) >= 30 and not replacing_primary:
            raise HTTPException(400, "每条申请最多30张证据截图。")
        if not isinstance(body.get("page_id"), str) or not body["page_id"]:
            raise HTTPException(400, "请选择截图页面。")
        page = one(conn, "file_pages", body["page_id"])
        if not page or not one(conn, "revision_files", revision_id=revision["id"], file_id=page["file_id"]):
            raise HTTPException(400, "截图页面不属于此版本附件。")
        kind = str(body.get("evidence_type", ""))
        if kind not in ("VALIDITY", "PARTICIPATION"):
            raise HTTPException(400, "请选择有效性证明或本人参与证明。")
        metadata = {"validity_kind": body.get("validity_kind"), "highlights": body.get("highlights", [])}
        if kind == "PARTICIPATION":
            attached_pages = [p for link in rows(conn, "revision_files", revision_id=revision["id"]) for p in rows(conn, "file_pages", file_id=link["file_id"])]
            matches = identity_matches(conn, attached_pages, context["student"], context["user"], context["class"]["name"])
            metadata["highlights"] = metadata["highlights"] or [{"box": match["box"], "field": match["field"]} for match in matches if match["page_id"] == page["id"]]
            selected = checked_box(body.get("box"), page["width_px"], page["height_px"])
            primary_box = quarter_page_box(page, selected[1])
            metadata["other_matches"] = [match for match in matches if match["page_id"] != page["id"] or _intersection(primary_box, match["box"]) != match["box"]]
        crop = create_crop(conn, revision, page, kind, body.get("box"), body.get("masks", []), metadata=metadata)
        _mark_manual_evidence(conn, revision)
        _touch(conn, context["application"])
        return crop_view(crop, conn=conn)


@router.delete("/applications/{application_id}/crops/{crop_id}")
def remove_crop(application_id: str, crop_id: str, user: dict = Depends(current_user)):
    with db.write() as conn:
        context, revision = _editable(conn, application_id, user)
        crop = one(conn, "evidence_crops", crop_id)
        if crop and crop["revision_id"] != revision["id"] and crop["revision_id"] == context.get("previous_revision_id"):
            # The first edit of a rejected frozen version clones attachments and crops.
            # Resolve the displayed old id to that exact clone, never delete history.
            crop = next((candidate for candidate in rows(conn, "evidence_crops", revision_id=revision["id"])
                         if all(candidate.get(key) == crop.get(key) for key in ("page_id", "evidence_type", "source_hash", "polygon_json", "sort_order"))), None)
        if not crop or crop["revision_id"] != revision["id"]:
            raise HTTPException(404, "截图不属于当前可编辑版本。")
        conn.execute(delete(tables["evidence_crops"]).where(tables["evidence_crops"].c.id == crop["id"]))
        _mark_manual_evidence(conn, revision)
        _touch(conn, context["application"])
    return {"ok": True}


@router.get("/crops/{crop_id}/image")
def get_crop(crop_id: str, user: dict = Depends(current_user)):
    with db.read() as conn:
        crop = one(conn, "evidence_crops", crop_id)
        if not crop:
            raise HTTPException(404, "截图不存在。")
        page = one(conn, "file_pages", crop["page_id"])
        _can_read_file(conn, page["file_id"], user)
    return _private_response(crop["private_asset_uri"])


@router.get("/crops/{crop_id}/public")
def get_public_crop(crop_id: str, user: dict = Depends(current_user)):
    with db.read() as conn:
        crop = one(conn, "evidence_crops", crop_id)
        if not crop or crop.get("public_asset_status") != "READY" or crop.get("redaction_version") != REDACTION_VERSION:
            raise HTTPException(404, "该截图尚未生成安全的公示版本。")
        revision = one(conn, "application_revisions", crop["revision_id"])
        context = _app_context(conn, revision["application_id"])
        current_round = one(conn, "submission_rounds", context["application"].get("current_round_id")) if context["application"].get("current_round_id") else None
        publicity = one(conn, "class_publicity", class_year_id=context["class_year"]["id"])
        if not publicity or not publicity["is_open"] or not current_round or current_round["state"] != "APPROVED" or current_round["revision_id"] != revision["id"]:
            raise HTTPException(403, "此截图当前未公开。")
        require(conn, user, "publicity.read", class_id=context["class"]["id"], owner_id=context["user"]["id"])
        if user["base_identity"] == "STUDENT":
            profile = one(conn, "student_profiles", user_id=user["id"])
            if not profile or not one(conn, "enrollments", student_id=profile["id"], class_year_id=context["class_year"]["id"], status="ACTIVE"):
                raise HTTPException(403, "只能查看本班本学年的公示截图。")
        page = one(conn, "file_pages", crop["page_id"])
        meta = crop.get("masks_json") or {}
        if meta.get("page_hash") != page["render_sha256"]:
            raise HTTPException(404, "公示截图来源校验失败。")
        if not asset_path(crop["private_asset_uri"]).is_file() or hashlib.sha256(asset_path(crop["private_asset_uri"]).read_bytes()).hexdigest() != crop["source_hash"]:
            raise HTTPException(404, "公示截图来源校验失败。")
    return _private_response(crop["public_asset_uri"])


def revision_fingerprint(conn, revision: dict) -> str:
    attachments = sorted(rows(conn, "revision_files", revision_id=revision["id"]), key=lambda r: (r["sort_order"], r["id"]))
    crops = sorted(rows(conn, "evidence_crops", revision_id=revision["id"]), key=lambda c: c["id"])
    body = {"policy_version": OCR_POLICY_VERSION, "revision_id": revision["id"], "title": revision["title"], "score": str(revision["requested_score"]), "activity_id": revision.get("activity_id"), "payload": revision.get("payload_json") or {}, "files": [link["file_id"] for link in attachments], "crops": [{"id": c["id"], "hash": c["source_hash"], "box": c["polygon_json"], "manual_masks": (c.get("masks_json") or {}).get("manual", [])} for c in crops]}
    body["max_automatic_pages"] = _page_limit()
    return hashlib.sha256(json.dumps(body, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def ocr_input_fingerprint(conn, revision: dict) -> str:
    """An OCR task depends on source pages, not a student's title, score or crops."""
    attachments = sorted(rows(conn, "revision_files", revision_id=revision["id"]), key=lambda item: (item["sort_order"], item["id"]))
    sources = []
    for attachment in attachments:
        file = one(conn, "file_objects", attachment["file_id"])
        pages = sorted(rows(conn, "file_pages", file_id=file["id"]), key=lambda page: (page["page_index"], page["id"]))
        sources.append({
            "file_id": file["id"], "sha256": file["sha256"], "byte_size": file["byte_size"],
            "mime_type": file["mime_type"], "sort_order": attachment["sort_order"],
            "pages": [{key: page[key] for key in ("id", "page_index", "render_version", "render_sha256", "width_px", "height_px", "rotation_degrees")} for page in pages],
        })
    return hashlib.sha256(json.dumps({"policy_version": OCR_POLICY_VERSION, "max_automatic_pages": _page_limit(), "revision_id": revision["id"], "sources": sources}, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def _usage(conn, action: str) -> int:
    month = now()[:7]
    table = tables["recognition_runs"]
    return conn.scalar(select(func.count()).select_from(table).where(and_(table.c.api_name == action, table.c.provider == "TENCENT", table.c.started_at >= month + "-01"))) or 0


@router.get("/integrations/status")
def integration_status(user: dict = Depends(current_user)):
    config = configuration()
    with db.read() as conn:
        general_used, seal_used = _usage(conn, "GeneralAccurateOCR"), _usage(conn, "SealOCR")
    return {"enabled": config["enabled"], "tencent_configured": bool(config["secret_id"] and config["secret_key"]), "deepseek_configured": bool(config["deepseek_key"]), "deepseek_model": config["deepseek_model"], "free_only": True, "general": {"used": general_used, "limit": config["general_limit"]}, "seal": {"used": seal_used, "limit": config["seal_limit"]}, "max_analysis_pages": config["max_analysis_pages"], "message": "仅调用已开通服务，不自动重试或购买资源；本地预算不代表云端共享余额。额度不足时可手动填写和截图。"}


def _ocr_page(page: dict, action: str, config: dict) -> dict:
    with OCR_CALL_LOCK:
        return _ocr_page_locked(page, action, config)


def _save_regions(conn, run: dict, data: dict):
    if run["api_name"] == "GeneralAccurateOCR":
        for index, item in enumerate(data.get("TextDetections") or []):
            polygon = item.get("Polygon") or _rect_polygon(item.get("ItemPolygon") or {})
            words = []
            for word, coords in zip(item.get("Words") or [], item.get("WordCoordPoint") or []):
                words.append({"text": str(word.get("Character", "")), "polygon": coords.get("WordCoordinate") or [], "confidence": word.get("Confidence")})
            add(conn, "recognition_regions", run_id=run["id"], region_no=index, kind="TEXT", text=str(item.get("DetectedText", "")), confidence=Decimal(str(item.get("Confidence") or 0)) / 100, polygon_json=polygon, extra_json={"words": words})
    else:
        seals = data.get("SealInfos") or ([data] if data.get("Location") else [])
        for index, item in enumerate(seals):
            add(conn, "recognition_regions", run_id=run["id"], region_no=index, kind="SEAL", text=str(item.get("SealBody", "")), confidence=None, polygon_json=_rect_polygon(item.get("Location") or {}), extra_json={"shape": item.get("SealShape"), "other_texts": item.get("OtherTexts", [])})


def _ocr_page_locked(page: dict, action: str, config: dict) -> dict:
    param_hash = hashlib.sha256((page["render_sha256"] + action + "v2-words").encode()).hexdigest()
    with db.write() as conn:
        cached = [r for r in rows(conn, "recognition_runs", api_name=action, parameters_hash=param_hash) if r["status"] == "SUCCEEDED"]
        if cached:
            own = [r for r in cached if r["page_id"] == page["id"]]
            if own:
                return own[-1]
            source = cached[-1]
            run = add(conn, "recognition_runs", page_id=page["id"], provider="CACHE", api_name=action, parameters_hash=param_hash, status="SUCCEEDED", provider_request_id=source.get("provider_request_id"), result_json=source["result_json"], started_at=None, finished_at=now())
            _save_regions(conn, run, source["result_json"] or {})
            return run
        limit = config["general_limit"] if action == "GeneralAccurateOCR" else config["seal_limit"]
        if _usage(conn, action) >= limit:
            raise ProviderError("LocalBudgetExceeded", "本月本地 OCR 保护预算已用完，已停止调用，可手动继续。")
        run = add(conn, "recognition_runs", page_id=page["id"], provider="TENCENT", api_name=action, parameters_hash=param_hash, status="RUNNING", provider_request_id=None, result_json=None, started_at=now(), finished_at=None)
    try:
        data = tencent_ocr(action, asset_path(page["image_uri"]).read_bytes(), config)
    except ProviderError as error:
        with db.write() as conn:
            change(conn, "recognition_runs", run["id"], status="FAILED", result_json={"error_code": error.code, "message": error.message}, finished_at=now())
        raise
    with db.write() as conn:
        run = change(conn, "recognition_runs", run["id"], status="SUCCEEDED", provider_request_id=data.get("RequestId"), result_json=data, finished_at=now())
        _save_regions(conn, run, data)
        return run


def _rect_polygon(rect: dict) -> list:
    return _polygon([int(rect.get("X", 0)), int(rect.get("Y", 0)), int(rect.get("Width", 0)), int(rect.get("Height", 0))])


def _compact(text) -> str:
    return re.sub(r"\s+", "", str(text or "")).casefold()


def _page_regions(conn, page: dict) -> list[dict]:
    # Latest successful result per API, not one copy of each cached execution.
    latest = {}
    for run in rows(conn, "recognition_runs", page_id=page["id"]):
        if run["status"] == "SUCCEEDED" and (run["api_name"] not in latest or (run.get("finished_at") or "") >= (latest[run["api_name"]].get("finished_at") or "")):
            latest[run["api_name"]] = run
    return [region for run in latest.values() for region in rows(conn, "recognition_regions", run_id=run["id"])]


def _automatic_regions(conn, page: dict, policy=None) -> list[dict]:
    """Current suggestion policy only; historical/public/manual reads remain intact."""
    _, cloud_allowed, boundaries = policy or _file_policy(conn, page["file_id"])
    native = _native_text_cached(conn, page)
    latest = {}
    for run in rows(conn, "recognition_runs", page_id=page["id"]):
        if run["status"] != "SUCCEEDED":
            continue
        action = run["api_name"]
        allowed = (action == NATIVE_TEXT_API and native and run["id"] == native["id"]) or (action == "GeneralAccurateOCR" and (cloud_allowed or page["id"] in boundaries) and not native) or (action == "SealOCR" and page["id"] in boundaries)
        if allowed and (action not in latest or (run.get("finished_at") or "") >= (latest[action].get("finished_at") or "")):
            latest[action] = run
    return [region for run in latest.values() for region in rows(conn, "recognition_regions", run_id=run["id"])]


def _native_pages(file: dict, pages: list[dict]) -> dict:
    """Persist local text with real character geometry; never spends OCR quota."""
    if file["mime_type"] != "application/pdf":
        return {}
    cached, missing = {}, []
    with db.read() as conn:
        for page in pages:
            parameter_hash = hashlib.sha256((file["sha256"] + str(page["page_index"]) + page["render_sha256"] + NATIVE_TEXT_API + "textlayer-v1").encode()).hexdigest()
            runs = rows(conn, "recognition_runs", page_id=page["id"], api_name=NATIVE_TEXT_API, parameters_hash=parameter_hash)
            if runs:
                cached[page["id"]] = sorted(runs, key=lambda item: item.get("finished_at") or "")[-1]
            else:
                missing.append((page, parameter_hash))
    if not missing:
        return cached
    from .pdf_text import extract_pages_text
    with PDF_LOCK:
        extracted_pages = extract_pages_text(asset_path(file["storage_uri"]), [page for page, _ in missing])
    for page, parameter_hash in missing:
        extracted = extracted_pages.get(page["id"], extracted_pages.get(page["page_index"], {}))
        usable = bool(extracted.get("usable") and str(extracted.get("text") or "").strip())
        # Large searchable PDFs retain every page, without holding the SQLite
        # writer lock across the entire document.
        with db.write() as conn:
            run = add(conn, "recognition_runs", page_id=page["id"], provider="LOCAL_PDF", api_name=NATIVE_TEXT_API, parameters_hash=parameter_hash, status="SUCCEEDED", provider_request_id=None, result_json={"usable": usable, "text": str(extracted.get("text") or ""), "reason": str(extracted.get("reason") or "")[:200]}, started_at=None, finished_at=now())
            if usable:
                for index, region in enumerate(extracted.get("regions") or []):
                    add(conn, "recognition_regions", run_id=run["id"], region_no=index, kind="TEXT", text=str(region.get("text") or ""), confidence=None, polygon_json=region.get("polygon_json") or [], extra_json=region.get("extra_json") or {})
            cached[page["id"]] = run
    return cached


def _token_box(region: dict, token: str) -> list[int]:
    words = (region.get("extra_json") or {}).get("words", [])
    text = "".join(_compact(word.get("text")) for word in words)
    start = text.find(token)
    if start < 0:
        return polygon_box(region["polygon_json"])
    cursor, boxes = 0, []
    for word in words:
        end = cursor + len(_compact(word.get("text")))
        if end > start and cursor < start + len(token) and word.get("polygon"):
            boxes.append(polygon_box(word["polygon"]))
        cursor = end
    if not boxes:
        return polygon_box(region["polygon_json"])
    left, top = min(box[0] for box in boxes), min(box[1] for box in boxes)
    right, bottom = max(box[0] + box[2] for box in boxes), max(box[1] + box[3] for box in boxes)
    return [left, top, right - left, bottom - top]


def identity_matches(conn, pages: list[dict], student: dict, student_user: dict, class_name: str, automatic=False) -> list[dict]:
    matches = []
    tokens = [("STUDENT_NUMBER", _compact(student.get("student_number"))), ("NAME", _compact(student_user.get("display_name"))), ("CLASS", _compact(class_name))]
    seen = set()
    max_pages = _page_limit() if automatic else None
    policies = {file_id: _file_policy(conn, file_id, max_pages) for file_id in {page["file_id"] for page in pages}} if automatic else {}
    for page in pages:
        file = one(conn, "file_objects", page["file_id"])
        native = _native_text_cached(conn, page)
        for region in (_automatic_regions(conn, page, policies[page["file_id"]]) if automatic else _page_regions(conn, page)):
            if region["kind"] != "TEXT":
                continue
            if automatic and not policies[page["file_id"]][1] and (not native or region["run_id"] != native["id"]):
                continue  # Long scanned files require the student to choose personal context.
            text = _compact(region.get("text"))
            for field, token in tokens:
                if not token or token not in text:
                    continue
                if field == "STUDENT_NUMBER" and token.isdigit() and not re.search(r"(?<![0-9])" + re.escape(token) + r"(?![0-9])", text):
                    continue
                box = _token_box(region, token)
                left, top = max(0, box[0]), max(0, box[1])
                right, bottom = min(page["width_px"], box[0] + box[2]), min(page["height_px"], box[1] + box[3])
                box = [left, top, right - left, bottom - top]
                if box[2] < 2 or box[3] < 2:
                    continue
                key = (page["id"], field, tuple(box))
                if key in seen:
                    continue
                seen.add(key)
                matches.append({"page_id": page["id"], "page_number": page["page_index"], "file_name": file["original_name"], "box": box, "field": field, "region_id": region["id"], "origin": "NATIVE_PDF_TEXT" if native and region["run_id"] == native["id"] else "TENCENT_OCR", "confidence": str(region.get("confidence") or 0), "url": f"/api/pages/{page['id']}/image"})
    order = {page["id"]: index for index, page in enumerate(pages)}
    priority = {"STUDENT_NUMBER": 0, "NAME": 1, "CLASS": 2}
    return sorted(matches, key=lambda match: (priority[match["field"]], -Decimal(match["confidence"]), order[match["page_id"]], match["box"][1]))


GENERIC_ACTIVITY_TITLES = {"活动", "证明", "活动证明", "参与证明", "集体活动", "集体活动证明", "社会服务", "社会服务证明", "思想政治理论学习", "思想政治理论学习证明", "文体竞赛", "奖励分", "奖励证明", "加分申请", "德育分申请", "学术科研及创新创业", "学生活动证明", "证明材料", "材料识别结果", "具体活动名称", "待填写活动名称", "特此证明", "学生组织社团任职", "培训", "培训证明", "讲座", "讲座证明", "任职证明", "社会实践", "社会实践证明", "竞赛证明", "活动参与", "活动参与证明", "加分小项", "未填写具体活动名称", "待分析", "未命名活动", "集体活动申报"}


def specific_activity_title(title: str) -> bool:
    text = re.sub(r"[\s，,。.!！:：;；()（）\[\]【】《》“”\"'\-—_]+", "", str(title or ""))
    if len(text) < 2 or text in GENERIC_ACTIVITY_TITLES:
        return False
    stripped = re.sub(r"(?:20\d{2}年?|[一二两三四五六七八九十百\d]+(?:次|场|项|个)|第[一二两三四五六七八九十百\d]+届|加分|申请|证明材料|证明|参加|参与)", "", text)
    return bool(stripped and stripped not in GENERIC_ACTIVITY_TITLES)


def _name_region(text: str, activity_names: list[str]) -> bool:
    compact = _compact(text)
    if not specific_activity_title(text):
        return False
    if any(_compact(name) in compact for name in activity_names if specific_activity_title(name)):
        return True
    if re.search(r"杯|广场舞|夜跑|运动会|党课|社会实践|志愿服务|讲座|培训|晚会|创新创业|互联网\+|挑战杯|任职|论坛", compact):
        return True
    labelled = re.search(r"(?:活动名称|项目名称|赛事名称)[：:]?(.+)", compact)
    return bool(labelled and specific_activity_title(labelled.group(1)))


def activity_title_supported(title: str, texts: list[dict], activities: list[dict]) -> bool:
    """Require some material anchor; this is not proof of authenticity or correctness."""
    normalize = lambda value: re.sub(r"[\W_]+", "", str(value or "")).casefold()
    source = normalize("".join(item.get("text", "") for item in texts))
    compact = normalize(title)
    if compact and compact in source:
        return True
    for activity in activities:
        name = normalize(activity["name"])
        if len(name) >= 2 and name in compact and name in source:
            return True
    for anchor in ("广场舞", "夜跑", "运动会", "蓝桥杯", "挑战杯", "青年大学习", "党课", "暑期社会实践", "互联网创新创业"):
        if normalize(anchor) in compact and normalize(anchor) in source:
            return True
    for fragment in re.split(r"[（()）:：，,《》“”\s]+", title):
        fragment = re.sub(r"20\d{2}年?|[一二两三四五六七八九十百\d]+次|证明|参与|参加", "", fragment)
        if len(fragment) >= 3 and specific_activity_title(fragment) and normalize(fragment) in source:
            return True
    return False


def _evidence_suggestions(conn, pages: list[dict], student: dict, student_user: dict, class_name: str, activity_names: list[str] | None = None) -> list[dict]:
    result = []
    last_seals = {}
    activity_names = activity_names or []
    max_pages = _page_limit()
    policies = {file_id: _file_policy(conn, file_id, max_pages) for file_id in {page["file_id"] for page in pages}}
    for page in pages:
        file = one(conn, "file_objects", page["file_id"])
        for region in _automatic_regions(conn, page, policies[page["file_id"]]):
            text = region.get("text") or ""
            kinds = []
            if region["kind"] == "SEAL":
                kinds.append("SEAL")
            if region["kind"] == "TEXT" and re.search(r"(?:19|20)\d{2}\s*(?:年|[./\-])\s*\d{1,2}|\d{1,2}\s*月\s*\d{1,2}\s*日", text):
                kinds.append("TIME")
            if region["kind"] == "TEXT" and _name_region(text, activity_names):
                kinds.append("NAME")
            for kind in kinds:
                if kind != "SEAL" and len([item for item in result if item.get("validity_kind") == kind]) >= 4:
                    continue
                x, y, w, h = polygon_box(region["polygon_json"])
                left, top = max(0, x - 10), max(0, y - 8)
                right, bottom = min(page["width_px"], x + w + 10), min(page["height_px"], y + h + 8)
                if right - left < 2 or bottom - top < 2:
                    continue
                suggestion = {"page_id": page["id"], "file_name": file["original_name"], "page_number": page["page_index"], "evidence_type": "VALIDITY", "validity_kind": kind, "box": [left, top, right - left, bottom - top], "region_id": region["id"]}
                if kind == "SEAL":
                    order = (page["page_index"], y, x)
                    if file["id"] not in last_seals or order > last_seals[file["id"]][0]:
                        last_seals[file["id"]] = (order, suggestion)
                else:
                    result.append(suggestion)
    result.extend(value[1] for value in last_seals.values())
    matches = identity_matches(conn, pages, student, student_user, class_name, automatic=True)
    if matches:
        primary = matches[0]
        page = next(p for p in pages if p["id"] == primary["page_id"])
        box = quarter_page_box(page, primary["box"][1] + primary["box"][3] / 2 - page["height_px"] / 8)
        highlights, other = [], []
        for match in matches:
            overlap = _intersection(box, match["box"]) if match["page_id"] == page["id"] else None
            if overlap:
                highlights.append({"box": match["box"], "field": match["field"]})
            if not overlap or overlap != match["box"]:
                other.append(match)
        result.append({"page_id": page["id"], "file_name": primary["file_name"], "page_number": page["page_index"], "evidence_type": "PARTICIPATION", "box": box, "region_id": primary["region_id"], "origin": primary["origin"], "highlights": highlights, "other_matches": other, "layout_version": "context-quarter-v2"})
    return result


def _last_seal_suggestions(conn, suggestions: list[dict]) -> list[dict]:
    """Normalize persisted V2 results too; provider array order never defines 'last'."""
    normalized, last = [], {}
    for index, item in enumerate(suggestions):
        item = dict(item)
        region = one(conn, "recognition_regions", item.get("region_id")) if item.get("region_id") else None
        candidate = {"source": "AI", "revision_id": None, "page_id": item.get("page_id"), "evidence_type": item.get("evidence_type"), "recognition_region_id": item.get("region_id"), "masks_json": {"evidence": {"validity_kind": item.get("validity_kind")}}}
        if not automatic_crop_policy(conn, candidate)[0]:
            continue
        if item.get("evidence_type") == "VALIDITY" and (item.get("validity_kind") == "SEAL" or (region and region["kind"] == "SEAL")):
            item["validity_kind"] = "SEAL"
            page = one(conn, "file_pages", item.get("page_id"))
            if page:
                box = polygon_box(region["polygon_json"]) if region else item.get("box", [0, 0, 0, 0])
                order = (page["page_index"], box[1], box[0])
                if page["file_id"] not in last or order > last[page["file_id"]][0]:
                    last[page["file_id"]] = (order, index)
        normalized.append((index, item))
    selected = {value[1] for value in last.values()}
    return [item for index, item in normalized if item.get("validity_kind") != "SEAL" or index in selected]


def _result_evidence(conn, context: dict, revision: dict, result: dict) -> list[dict]:
    """Re-derive existing successful OCR, including old jobs, without making any cloud call."""
    pages = []
    for page_id in result.get("page_ids", []):
        page = one(conn, "file_pages", page_id)
        if page and one(conn, "revision_files", revision_id=revision["id"], file_id=page["file_id"]):
            pages.append(page)
    if pages and any(_page_regions(conn, page) for page in pages):
        names = [activity["canonical_name"] for activity in rows(conn, "activities", status="ACTIVE")]
        if specific_activity_title(revision.get("title", "")):
            names.append(revision["title"])
        return _evidence_suggestions(conn, pages, context["student"], context["user"], context["class"]["name"], names)
    return _last_seal_suggestions(conn, result.get("suggested_crops", []))


@router.post("/applications/{application_id}/analyze")
def analyze(application_id: str, body: dict | None = None, user: dict = Depends(current_user)):
    if not ANALYSIS_LOCK.acquire(blocking=False):
        raise HTTPException(409, "已有识别任务正在进行，请稍后重试。")
    try:
        return run_ai(application_id, user)
    finally:
        ANALYSIS_LOCK.release()


def _recognition_snapshot(conn, application_id: str, user: dict, config: dict):
    from .submissions import assert_confirmed
    context, revision = _editable(conn, application_id, user)
    assert_confirmed(conn, application_id)
    links = sorted(rows(conn, "revision_files", revision_id=revision["id"]), key=lambda link: (link["sort_order"], link["id"]))
    if not links:
        raise HTTPException(400, "请先上传至少一份证明材料。")
    pages = [page for link in links for page in sorted(rows(conn, "file_pages", file_id=link["file_id"]), key=lambda p: p["page_index"])]
    return context, revision, pages, revision_fingerprint(conn, revision)


def run_ocr(application_id: str, user: dict, progress=None, _config: dict | None = None) -> dict:
    """Server-worker entry point: OCR only, never calls the language model."""
    config = _config or configuration()
    max_pages = _page_limit(config)
    with db.write() as conn:
        context, revision, pages, fingerprint = _recognition_snapshot(conn, application_id, user, config)
        source_fingerprint = ocr_input_fingerprint(conn, revision)
        files = {page["file_id"]: one(conn, "file_objects", page["file_id"]) for page in pages}
    common = {"policy_version": OCR_POLICY_VERSION, "application_id": application_id, "input_revision_id": revision["id"], "input_fingerprint": fingerprint, "ocr_fingerprint": source_fingerprint, "page_ids": [p["id"] for p in pages]}
    warnings, ocr_runs, tasks, file_results = [], [], [], []
    if progress:
        progress("检查PDF已有文字层", 0, 0)
    for file_id, file in files.items():
        file_pages = [page for page in pages if page["file_id"] == file_id]
        native = _native_pages(file, file_pages)
        native_ids = {page_id for page_id, run in native.items() if (run.get("result_json") or {}).get("usable")}
        if native_ids and progress:
            progress("读取现成文字", 0, 0)
        ocr_runs.extend(native[page["id"]] for page in file_pages if page["id"] in native_ids)
        cloud_allowed = len(file_pages) <= max_pages
        boundaries = {file_pages[0]["id"], file_pages[-1]["id"]}
        missing_native = [page["page_index"] for page in file_pages if page["id"] not in native_ids and page["id"] not in boundaries]
        file_results.append({"file_id": file_id, "file_name": file["original_name"], "page_count": len(file_pages), "max_pages": max_pages, "cloud_ocr_eligible": True,
                             "full_page_cloud_ocr_eligible": cloud_allowed, "cloud_ocr_scope": "FULL" if cloud_allowed else "ENDPOINTS", "participation_mode": "AUTO" if cloud_allowed or native_ids else "MANUAL",
                             "text_layer_check_complete": True, "text_layer_checked_page_count": len(file_pages) if file["mime_type"] == "application/pdf" else 0,
                             "native_text_status": "USED" if len(native_ids) == len(file_pages) else ("PARTIAL" if native_ids else "NONE"),
                             "text_layer_used": bool(native_ids), "native_pages": [page["page_index"] for page in file_pages if page["id"] in native_ids],
                             "manual_page_numbers": missing_native if not cloud_allowed else [], "seal_pages": sorted({file_pages[0]["page_index"], file_pages[-1]["page_index"]}), "warnings": []})
        for page in file_pages:
            if (cloud_allowed or page["id"] in boundaries) and page["id"] not in native_ids:
                tasks.append((page, "GeneralAccurateOCR"))
            if page["id"] in boundaries:
                tasks.append((page, "SealOCR"))
    done, total, stopped, stop_reason = 0, len(tasks), False, None
    if progress:
        progress("OCR准备", done, total)
    failed_tasks = []
    for page, action in tasks:
        # Disabling cloud access still permits explicitly cached results and local PDF text.
        if not config["enabled"] or not config["secret_id"] or not config["secret_key"]:
            param_hash = hashlib.sha256((page["render_sha256"] + action + "v2-words").encode()).hexdigest()
            with db.read() as conn:
                cached = any(run["status"] == "SUCCEEDED" for run in rows(conn, "recognition_runs", api_name=action, parameters_hash=param_hash))
            if not cached:
                stopped, stop_reason = True, "NotConfigured"
                warnings.append("腾讯云 OCR 未配置或已关闭，未继续云调用；可手动填写和框选。")
                break
        try:
            ocr_runs.append(_ocr_page(page, action, config))
        except ProviderError as error:
            warnings.append(error.message)
            failed_tasks.append((page["id"], action))
            # Stop all further cloud actions, including the other API, on quota/resource errors.
            if error.code in ("LocalBudgetExceeded", "NotConfigured", "NetworkError", "HTTPError") or any(part in error.code for part in ("AuthFailure", "Unauthorized", "ResourcePackageRunOut", "InArrears", "ChargeStatusException", "LimitExceeded", "ResourceUnavailable", "ResourcesSoldOut", "Quota", "UnOpenError")):
                stopped, stop_reason = True, error.code
        done += 1
        if progress:
            progress("识别文字" if action == "GeneralAccurateOCR" else "识别印章", done, total)
        if stopped:
            break
    with db.write() as conn:
        text_pages, cloud_text_pages, seal_pages = set(), set(), set()
        for run in ocr_runs:
            regions = rows(conn, "recognition_regions", run_id=run["id"])
            if run["api_name"] in ("GeneralAccurateOCR", NATIVE_TEXT_API) and any((region.get("text") or "").strip() for region in regions):
                text_pages.add(run["page_id"])
                if run["api_name"] == "GeneralAccurateOCR":
                    cloud_text_pages.add(run["page_id"])
            if run["api_name"] == "SealOCR":
                seal_pages.add(run["page_id"])
            page = one(conn, "file_pages", run["page_id"])
            for crop in rows(conn, "evidence_crops", page_id=page["id"], revision_id=revision["id"]):
                current_revision = one(conn, "application_revisions", crop["revision_id"])
                if current_revision.get("content_frozen_at"):
                    continue
                manual = (crop.get("masks_json") or {}).get("manual", [])
                change(conn, "evidence_crops", crop["id"], **_build_public_crop(conn, crop, page, manual))
        activities = rows(conn, "activities", status="ACTIVE")
        names = [activity["canonical_name"] for activity in activities]
        if specific_activity_title(revision.get("title", "")):
            names.append(revision["title"])
        evidence = _evidence_suggestions(conn, pages, context["student"], context["user"], context["class"]["name"], names)
        seal_files = {one(conn, "file_pages", item["page_id"])["file_id"] for item in evidence if item.get("validity_kind") == "SEAL"}
    skipped_files, per_file_warnings = [], []
    for item in file_results:
        file_pages = [page for page in pages if page["file_id"] == item["file_id"]]
        completed = sum(page["id"] in text_pages for page in file_pages)
        item["processed_page_count"] = completed
        item["text_covered_page_count"] = completed
        item["ocr_page_numbers"] = [page["page_index"] for page in file_pages if page["id"] in cloud_text_pages]
        item["automatic_ocr_status"] = "COMPLETED" if completed == len(file_pages) else ("PARTIAL" if completed else "MANUAL_REQUIRED")
        item["seal_found"] = item["file_id"] in seal_files
        boundary_ids = {file_pages[0]["id"], file_pages[-1]["id"]}
        item["seal_status"] = "FOUND" if item["seal_found"] else ("NOT_FOUND" if boundary_ids <= seal_pages else "LIMITED")
        item["seal_message"] = "已定位首页/末页中的最后一枚印章，请核对原件。" if item["seal_found"] else "仅检查文件首页和末页，未定位到印章；请自行寻找并手动框选。"
        if not item["seal_found"]:
            item["warnings"].append(item["seal_message"])
            per_file_warnings.append({"file_id": item["file_id"], "file_name": item["file_name"], "code": "SEAL_NOT_FOUND" if item["seal_status"] == "NOT_FOUND" else "SEAL_CHECK_LIMITED", "message": item["seal_message"]})
        if item["manual_page_numbers"]:
            reason = f"文件超过{max_pages}页：云OCR仅识别首页和末页；中间无文字层的扫描页不识别，请手动寻找并框选本人参与信息。"
            skipped_files.append({"file_id": item["file_id"], "file_name": item["file_name"], "page_count": item["page_count"], "max_pages": max_pages, "page_numbers": item["manual_page_numbers"], "reason": reason, "code": "MIDDLE_SCAN_PAGES_MANUAL"})
            item["warnings"].append(reason)
            per_file_warnings.append({"file_id": item["file_id"], "file_name": item["file_name"], "code": "MIDDLE_SCAN_PAGES_MANUAL", "message": reason})
        elif completed < len(file_pages):
            message = "部分页面未取得可用文字，请按原件手动补充；已成功结果予以保留。"
            item["warnings"].append(message)
            per_file_warnings.append({"file_id": item["file_id"], "file_name": item["file_name"], "code": "TEXT_INCOMPLETE", "message": message})
    warnings.extend(warning["message"] for warning in per_file_warnings)
    partial = bool(stopped or failed_tasks or skipped_files or len(text_pages) < len(pages))
    if partial:
        warnings.append("本次仅完成部分自动识别，未覆盖页面或印章需人工核对，已成功结果不会丢失。")
    found_validity = {item["validity_kind"] for item in evidence if item["evidence_type"] == "VALIDITY"}
    missing = [kind for kind in ("NAME", "TIME", "SEAL") if kind not in found_validity]
    for kind, label in (("NAME", "具体活动名称"), ("TIME", "活动时间"), ("SEAL", "印章")):
        if kind in missing:
            warnings.append(f"未定位到{label}，请人工核对原件。")
    if not any(item["evidence_type"] == "PARTICIPATION" for item in evidence):
        warnings.append("未定位到账户对应的姓名、班级或学号，请人工核对并选择主截图。")
    return {**common, "status": "SUCCEEDED" if text_pages else "MANUAL_REQUIRED", "partial": partial, "stopped": stopped, "stop_reason": stop_reason,
            "planned_tasks": total, "completed_tasks": done, "native_text_page_count": sum(len(item["native_pages"]) for item in file_results),
            "skipped_files": skipped_files, "manual_required_files": skipped_files, "file_results": file_results, "processed_files": [item for item in file_results if item["processed_page_count"]], "processed_page_ids": [page["id"] for page in pages if page["id"] in text_pages], "per_file_warnings": per_file_warnings,
            "run_ids": list(dict.fromkeys(run["id"] for run in ocr_runs)), "suggested_crops": evidence, "missing_validity": missing, "warnings": list(dict.fromkeys(warnings)), "message": ("已取得可用文字，部分页面或印章仍需人工核对。" if partial else "文字识别已完成，请核对原页定位；识别结果不作为材料真实性结论。") if text_pages else "未取得可供分析的文字，可手动填写、计分和框选证明。"}


def apply_ocr_evidence(application_id: str, user: dict, result: dict, expected_fingerprint=None, on_applied=None, expected_ocr_fingerprint=None) -> dict:
    """Apply evidence only after the worker rechecks the exact immutable input snapshot."""
    from .submissions import assert_confirmed
    with db.write() as conn:
        context, revision = _editable(conn, application_id, user)
        assert_confirmed(conn, application_id)
        fingerprint = revision_fingerprint(conn, revision)
        expected = expected_fingerprint or result.get("input_fingerprint")
        if expected_ocr_fingerprint is not None:
            valid_input = bool(expected_ocr_fingerprint and result.get("ocr_fingerprint") == expected_ocr_fingerprint and ocr_input_fingerprint(conn, revision) == expected_ocr_fingerprint)
        else:
            valid_input = fingerprint == expected
        if result.get("input_revision_id") != revision["id"] or not valid_input:
            raise HTTPException(409, "OCR期间材料或内容已变动，建议保留但不会覆盖现稿。")
        existing = rows(conn, "evidence_crops", revision_id=revision["id"])
        if (revision.get("payload_json") or {}).get("evidence_manually_edited") or any(crop["source"] == "STUDENT" for crop in existing):
            if on_applied:
                on_applied(conn, fingerprint, fingerprint)
            return {"fingerprint": fingerprint, "input_revision_id": revision["id"], "applied": False, "message": "保留学生已手动调整的截图。"}
        for crop in existing:
            if crop["source"] == "AI":
                conn.execute(delete(tables["evidence_crops"]).where(tables["evidence_crops"].c.id == crop["id"]))
        _apply_evidence(conn, revision, _result_evidence(conn, context, revision, result))
        _touch(conn, context["application"])
        updated = one(conn, "application_revisions", revision["id"])
        after_fingerprint = revision_fingerprint(conn, updated)
        if on_applied:
            on_applied(conn, fingerprint, after_fingerprint)
        return {"fingerprint": after_fingerprint, "before_fingerprint": fingerprint, "input_revision_id": revision["id"], "applied": True}


def run_ai(application_id: str, user: dict, progress=None, ocr_result=None) -> dict:
    return _analyze(application_id, user, configuration(), progress, ocr_result=ocr_result)


def _mark_manual_evidence(conn, revision: dict):
    payload = dict(revision.get("payload_json") or {})
    payload["evidence_manually_edited"] = True
    change(conn, "application_revisions", revision["id"], payload_json=payload)


def _apply_evidence(conn, revision: dict, suggestions: list) -> None:
    if revision.get("content_frozen_at"):
        return
    suggestions = _last_seal_suggestions(conn, suggestions)
    existing = rows(conn, "evidence_crops", revision_id=revision["id"])
    manual_evidence = (revision.get("payload_json") or {}).get("evidence_manually_edited")
    if not manual_evidence:
        desired_seals = {}
        for suggestion in suggestions:
            if suggestion.get("validity_kind") == "SEAL":
                page = one(conn, "file_pages", suggestion["page_id"])
                if page:
                    desired_seals[page["file_id"]] = (page["id"], suggestion["box"])
        for crop in list(existing):
            if crop["source"] != "AI":
                continue
            if not automatic_crop_policy(conn, crop, revision)[0]:
                conn.execute(delete(tables["evidence_crops"]).where(tables["evidence_crops"].c.id == crop["id"]))
                existing.remove(crop)
                continue
            metadata = (crop.get("masks_json") or {}).get("evidence") or {}
            region = one(conn, "recognition_regions", crop["recognition_region_id"]) if crop.get("recognition_region_id") else None
            if metadata.get("validity_kind") != "SEAL" and not (region and region["kind"] == "SEAL"):
                continue
            page = one(conn, "file_pages", crop["page_id"])
            wanted = desired_seals.get(page["file_id"])
            if not wanted or (page["id"], polygon_box(crop["polygon_json"])) != wanted:
                conn.execute(delete(tables["evidence_crops"]).where(tables["evidence_crops"].c.id == crop["id"]))
                existing.remove(crop)
    for item in sorted(suggestions, key=lambda item: item.get("evidence_type") != "PARTICIPATION")[:25]:
        if manual_evidence:
            break  # Explicitly preserve a student's additions, removals and adjustments.
        page = one(conn, "file_pages", item.get("page_id"))
        if not page or not one(conn, "revision_files", revision_id=revision["id"], file_id=page["file_id"]):
            continue
        if len(existing) >= 30:
            break
        if any(crop["page_id"] == page["id"] and crop["evidence_type"] == item["evidence_type"] and polygon_box(crop["polygon_json"]) == item["box"] and (crop.get("masks_json") or {}).get("evidence", {}).get("validity_kind") == item.get("validity_kind") for crop in existing):
            continue
        metadata = {key: item[key] for key in ("file_name", "page_number", "highlights", "other_matches", "validity_kind", "layout_version", "origin") if key in item}
        crop = create_crop(conn, revision, page, item["evidence_type"], item["box"], source="AI", region_id=item.get("region_id"), metadata=metadata)
        existing.append(crop)


def _ai_material_excerpt(conn, ocr: dict, student: dict, student_user: dict, max_chars=48000):
    """Bound the LLM input, while full local text/geometry remains in recognition storage."""
    records, seen, policies = [], set(), {}
    max_pages = _page_limit()
    for run_id in ocr.get("run_ids", []):
        run = one(conn, "recognition_runs", run_id)
        if not run or run["status"] != "SUCCEEDED" or run["id"] in seen:
            continue
        seen.add(run["id"])
        page = one(conn, "file_pages", run["page_id"])
        if page["file_id"] not in policies:
            policies[page["file_id"]] = _file_policy(conn, page["file_id"], max_pages)
        _, cloud_allowed, boundaries = policies[page["file_id"]]
        if (run["api_name"] == "GeneralAccurateOCR" and ((not cloud_allowed and page["id"] not in boundaries) or _native_text_cached(conn, page))) or (run["api_name"] == "SealOCR" and page["id"] not in boundaries):
            continue
        text = (run.get("result_json") or {}).get("text") if run["api_name"] == NATIVE_TEXT_API else None
        if text is None:
            text = "\n".join(region.get("text") or "" for region in rows(conn, "recognition_regions", run_id=run_id))
        if str(text).strip():
            records.append({"page_id": page["id"], "file_id": page["file_id"], "page_number": page["page_index"], "kind": run["api_name"], "text": str(text)})
    source_chars = sum(len(record["text"]) for record in records)
    source_page_count = len({record["page_id"] for record in records})
    selected = []
    if source_chars <= max_chars:
        selected = records
    else:
        tokens = [str(value).strip() for value in (student.get("student_number"), student_user.get("display_name")) if str(value or "").strip()]
        patterns = [re.compile((r"(?<![0-9])" if token.isdigit() else "") + r"\s*".join(re.escape(character) for character in token) + (r"(?![0-9])" if token.isdigit() else "")) for token in tokens]
        boundaries = {}
        for file_id in {record["file_id"] for record in records}:
            file_pages = rows(conn, "file_pages", file_id=file_id)
            boundaries[file_id] = {min(page["page_index"] for page in file_pages), max(page["page_index"] for page in file_pages)}
        matched, match_pages = [], set()
        for record in records:
            spans = [match.span() for pattern in patterns for match in pattern.finditer(record["text"])][:12]
            if spans:
                windows = []
                for start, end in sorted(spans):
                    left, right = max(0, start - 650), min(len(record["text"]), end + 650)
                    if windows and left <= windows[-1][1]:
                        windows[-1] = (windows[-1][0], max(right, windows[-1][1]))
                    else:
                        windows.append((left, right))
                text = "\n[该页其余内容略]\n".join(record["text"][left:right] for left, right in windows)[:8000]
                matched.append({**record, "text": text, "excerpt_reason": "账户姓名或学号命中上下文", "page_excerpt": len(text) < len(record["text"])})
                match_pages.add((record["file_id"], record["page_number"]))
        # Reserve space for file boundaries, so a very long roster cannot crowd out
        # the activity heading/date. Match windows still receive most of the budget.
        budget, selected_keys = 32000, set()
        for record in matched:
            if budget <= 0:
                break
            text = record["text"][:budget]
            selected.append({**record, "text": text})
            selected_keys.add((record["page_id"], record["kind"]))
            budget -= len(text)
        remaining = max_chars - sum(len(record["text"]) for record in selected)
        candidates = []
        for record in records:
            if (record["page_id"], record["kind"]) in selected_keys:
                continue
            if record["page_number"] in boundaries[record["file_id"]]:
                candidates.append((0, record))
            elif any(file_id == record["file_id"] and abs(number - record["page_number"]) == 1 for file_id, number in match_pages):
                candidates.append((1, record))
        for priority, record in sorted(candidates, key=lambda pair: (pair[0], pair[1]["file_id"], pair[1]["page_number"])):
            if remaining <= 0:
                break
            text = record["text"][:min(3000 if priority == 0 else 1200, remaining)]
            selected.append({**record, "text": text, "excerpt_reason": "原文件首页或末页" if priority == 0 else "本人命中页的相邻页", "page_excerpt": len(text) < len(record["text"])})
            remaining -= len(text)
    selected_chars = sum(len(record["text"]) for record in selected)
    selected_pages = len({record["page_id"] for record in selected})
    excerpt = {"character_limit": max_chars, "source_character_count": source_chars, "selected_character_count": selected_chars,
               "source_page_count": source_page_count, "selected_page_count": selected_pages, "omitted_page_count": source_page_count - selected_pages,
               "omitted_character_count": max(0, source_chars - selected_chars), "coverage_complete": source_chars <= max_chars,
               "note": "全文已在本地保存；AI仅收到按本人姓名/学号命中、文件首页末页及邻页选择的摘录，未覆盖内容需核对原件。" if source_chars > max_chars else "AI收到已提取的可用文字；这不代表原件全部页面或印章均已自动识别。"}
    return selected, excerpt


def _analyze(application_id: str, user: dict, config: dict, progress=None, ocr_result=None) -> dict:
    ocr = ocr_result if ocr_result is not None else run_ocr(application_id, user, progress, _config=config)
    if ocr["status"] != "SUCCEEDED":
        return ocr
    with db.write() as conn:
        context, revision = _editable(conn, application_id, user)
        if revision["id"] != ocr["input_revision_id"] or revision_fingerprint(conn, revision) != ocr["input_fingerprint"]:
            raise HTTPException(409, "识别期间填写内容或附件已变更，本次不再生成覆盖现稿的AI结果。")
        if ocr.get("application_id", application_id) != application_id or (ocr.get("ocr_fingerprint") and ocr["ocr_fingerprint"] != ocr_input_fingerprint(conn, revision)):
            raise HTTPException(409, "OCR结果对应的材料来源已变更，请重新识别。")
        if not config["enabled"]:
            return {"status": "MANUAL_REQUIRED", "message": "云识别已在本地配置关闭，可手动填写和截图。"}
        pages = [one(conn, "file_pages", page_id) for page_id in ocr["page_ids"]]
        rule = one(conn, "score_rule_versions", context["class_year"]["score_rule_version_id"])
        analysis = add(conn, "ai_analyses", application_id=application_id, input_revision_id=revision["id"], input_fingerprint=ocr["input_fingerprint"], score_rule_version_id=rule["id"], model_name=config["deepseek_model"], prompt_version="moral-score-v2", status="RUNNING", suggested_title=None, suggested_score=None, suggested_activity_id=None, result_json=None)
        if ocr.get("policy_version") not in (None, OCR_POLICY_VERSION):
            raise HTTPException(409, "OCR结果使用旧识别策略，请按当前策略重新排队。")
        texts, excerpt = _ai_material_excerpt(conn, ocr, context["student"], context["user"])
        for index, page in enumerate(pages):
            add(conn, "analysis_inputs", analysis_id=analysis["id"], file_id=page["file_id"], page_id=page["id"], recognition_run_id=None, input_kind="PAGE", input_hash=page["render_sha256"], sort_order=index)
        for index, run_id in enumerate(ocr["run_ids"]):
            run = one(conn, "recognition_runs", run_id)
            page = one(conn, "file_pages", run["page_id"])
            add(conn, "analysis_inputs", analysis_id=analysis["id"], file_id=page["file_id"], page_id=page["id"], recognition_run_id=run_id, input_kind="OCR", input_hash=run["parameters_hash"], sort_order=len(pages) + index)
        activities = [{"id": a["id"], "name": a["canonical_name"], "aliases": [v["alias"] for v in rows(conn, "activity_aliases", activity_id=a["id"]) if v["status"] == "ACTIVE"]} for a in rows(conn, "activities", status="ACTIVE")][:300]
    evidence, warnings = ocr["suggested_crops"], list(ocr["warnings"])
    metadata = {key: ocr[key] for key in ("partial", "skipped_files", "manual_required_files", "file_results", "per_file_warnings", "planned_tasks", "completed_tasks", "native_text_page_count", "stopped", "stop_reason") if key in ocr}
    metadata.update(policy_version=OCR_POLICY_VERSION, context_excerpt=excerpt)
    if not excerpt["coverage_complete"]:
        warnings.append(excerpt["note"])
    if progress:
        progress("AI整理具体活动与分值", 0, 1)
    prompt_context = {"category": context["category"]["name"], "category_cap": str(context["category"]["default_cap"]), "academic_year": context["year"]["name"], "student": {"name": context["user"]["display_name"], "class": context["class"]["name"]}, "scoring_rules": rule.get("rules_json") or {}, "activity_catalog": activities, "context_excerpt": excerpt, "untrusted_materials": texts}
    try:
        suggestion = deepseek_suggestion(prompt_context, config)
        title = str(suggestion.get("title", "")).strip()[:255]
        score = money(suggestion.get("requested_score", "0"))
        if not specific_activity_title(title):
            raise ProviderError("ActivityNameMissing", "未识别出具体活动名称，请学生按原件补充，不能只填泛称‘活动证明’。")
        if not activity_title_supported(title, texts, activities):
            raise ProviderError("ActivityNameUncertain", "建议活动名未能在识别文字中找到足够对应，请按原件人工补全。")
        if score < 0 or score > 105:
            raise ValueError("invalid suggestion")
        with db.read() as conn:
            evidence = _evidence_suggestions(conn, pages, context["student"], context["user"], context["class"]["name"], [activity["name"] for activity in activities] + [title])
        valid_activity_ids = {activity["id"] for activity in activities}
        activity_id = suggestion.get("activity_id")
        if activity_id not in valid_activity_ids:
            activity_id = None
        payload = suggestion.get("payload_json") if isinstance(suggestion.get("payload_json"), dict) else {}
        payload = {key: str(payload.get(key, ""))[:500] for key in ("edition_key", "track_key", "group_key", "work_key", "award_level", "subtype")}
        suggestion = {"title": title, "requested_score": str(score), "score": str(score), "suggested_score": str(score), "activity_id": activity_id, "payload_json": payload, "explanation": str(suggestion.get("explanation", ""))[:3000], "match_uncertain": bool(suggestion.get("match_uncertain", not activity_id)), "suggested_crops": evidence}
    except (ProviderError, InvalidOperation, ValueError, TypeError, HTTPException) as error:
        message = error.message if isinstance(error, ProviderError) else "AI 建议不符合预期格式，请手动填写。"
        error_code = error.code if isinstance(error, ProviderError) else "InvalidSuggestion"
        with db.write() as conn:
            change(conn, "ai_analyses", analysis["id"], status="FAILED", result_json={**metadata, "error_code": error_code, "message": message, "warnings": list(dict.fromkeys(warnings)), "suggested_crops": evidence})
        return {**metadata, "status": "MANUAL_REQUIRED", "error_code": error_code, "analysis_id": analysis["id"], "analysis": {"id": analysis["id"], "status": "FAILED"}, "message": message, "warnings": list(dict.fromkeys(warnings)), "suggested_crops": evidence}
    with db.write() as conn:
        change(conn, "ai_analyses", analysis["id"], status="SUCCEEDED", suggested_title=suggestion["title"], suggested_score=score, suggested_activity_id=activity_id, result_json={**metadata, "suggestion": suggestion, "warnings": list(dict.fromkeys(warnings))})
    if progress:
        progress("AI分析完成", 1, 1)
    return {**metadata, "status": "SUCCEEDED", "input_revision_id": revision["id"], "input_fingerprint": ocr["input_fingerprint"], "analysis_id": analysis["id"], "analysis": {"id": analysis["id"], "status": "SUCCEEDED", "input_revision_id": revision["id"]}, "suggestion": suggestion, "warnings": list(dict.fromkeys(warnings)), "message": "建议已生成；当前填写内容尚未改变。"}


@router.post("/applications/{application_id}/apply-analysis")
def apply_analysis(application_id: str, body: dict, user: dict = Depends(current_user)):
    from .workflow import application_detail
    with db.write() as conn:
        context, revision = _editable(conn, application_id, user)
        if not isinstance(body.get("analysis_id"), str) or not body["analysis_id"]:
            raise HTTPException(400, "请选择分析记录。")
        analysis = one(conn, "ai_analyses", body["analysis_id"])
        if not analysis or analysis["application_id"] != application_id or analysis["status"] != "SUCCEEDED":
            raise HTTPException(400, "该分析记录不能应用。")
        if analysis["input_revision_id"] != revision["id"] or analysis["input_fingerprint"] != revision_fingerprint(conn, revision):
            raise HTTPException(409, "分析后内容或附件已变更，请重新分析；已有填写不会被覆盖。")
        suggestion = (analysis.get("result_json") or {}).get("suggestion", {})
        payload = {**(revision.get("payload_json") or {}), **suggestion.get("payload_json", {})}
        payload["ai_match_uncertain"] = suggestion.get("match_uncertain", False)
        payload["ai_explanation"] = suggestion.get("explanation", "")
        revision = change(conn, "application_revisions", revision["id"], title=suggestion["title"], requested_score=money(suggestion["requested_score"]), activity_id=suggestion.get("activity_id"), payload_json=payload, ai_analysis_id=analysis["id"])
        _apply_evidence(conn, revision, suggestion.get("suggested_crops", []))
        _touch(conn, context["application"])
        return application_detail(conn, one(conn, "applications", application_id), user)


def excel_text(value) -> str:
    value = str(value or "").replace("\x00", "")
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else value


def export_workbook(scores: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "德育分汇总"
    first = {1: "序号", 2: "班级", 3: "姓名", 4: "学号", 5: "品德素质分（80分）", 10: "能力拓展分（20分）", 13: "奖励分（5分）", 14: "总分", 15: "签字确认"}
    second = {5: "基准分（70分）", 6: "集体活动分（3分）", 7: "思想政治理论学习分（3分）", 8: "社会服务分（4分）", 9: "扣分（按规定）", 10: "学术科研及创新创业加分（10分）", 11: "文体竞赛加分（6分）", 12: "学生组织、社团任职加分（4分）"}
    for col, title in first.items():
        sheet.cell(1, col, title)
    for col, title in second.items():
        sheet.cell(2, col, title)
    for col in (1, 2, 3, 4, 13, 14, 15):
        sheet.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    sheet.merge_cells("E1:I1")
    sheet.merge_cells("J1:L1")
    for row_no, student in enumerate(scores["members"], start=3):
        sheet.cell(row_no, 1, row_no - 2)
        sheet.cell(row_no, 2, excel_text(scores["class"]["name"]))
        sheet.cell(row_no, 3, excel_text(student["display_name"]))
        sheet.cell(row_no, 4, excel_text(student.get("student_number", "")))
        sheet.cell(row_no, 5, 70)
        sheet.cell(row_no, 9, 0)
        categories = sorted(student["categories"], key=lambda c: c.get("sort_order", 0))
        category_cols = [6, 7, 8, 10, 11, 12, 13]
        max_lines = 1
        for category, col in zip(categories, category_cols):
            lines = [f"{i + 1}、{item.get('title', '')}（{money(item.get('score', item.get('requested_score', 0)))}分）" + ("〔同项目/兼任职务取最高，本项不重复计入〕" if item.get("counted") is False else "") for i, item in enumerate(category.get("items", []))]
            lines.append(f"合计：{money(category['score'])}分")
            max_lines = max(max_lines, len(lines))
            sheet.cell(row_no, col, excel_text("\n".join(lines)))
        sheet.cell(row_no, 14, float(money(student["total"])))
        sheet.cell(row_no, 15, "")
        sheet.row_dimensions[row_no].height = min(300, max(42, max_lines * 18))
    thin = Side(style="thin", color="B3BEC9")
    for row in sheet.iter_rows(min_row=1, max_row=max(2, sheet.max_row), max_col=15):
        for cell in row:
            cell.font = Font(name="宋体", size=10, bold=cell.row <= 2, color="FFFFFF" if cell.row <= 2 else "1F2937")
            cell.fill = PatternFill("solid", fgColor="173F66" if cell.row <= 2 else ("F2F6FA" if cell.row % 2 == 0 else "FFFFFF"))
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center" if cell.column in (1, 2, 3, 4, 5, 9, 14, 15) or cell.row <= 2 else "left", vertical="center", wrap_text=True)
            if cell.column in (5, 9, 14) and cell.row >= 3:
                cell.number_format = "0.00"
    widths = [7, 14, 12, 19, 12, 34, 34, 34, 12, 37, 34, 34, 30, 12, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 48
    sheet.freeze_panes = "F3"
    sheet.auto_filter.ref = f"A2:O{max(2, sheet.max_row)}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"
    sheet.print_area = f"A1:O{max(2, sheet.max_row)}"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@router.get("/export")
def export_scores(class_id: str, year_id: str, user: dict = Depends(current_user)):
    from urllib.parse import quote
    from .workflow import class_scores
    with db.read() as conn:
        require(conn, user, "export.read", class_id=class_id)
        scores = class_scores(conn, class_id, year_id, public=False)
    filename = f"{scores['year']['name']}-{scores['class']['name']}-德育分汇总.xlsx"
    content = export_workbook(scores)
    return StreamingResponse(io.BytesIO(content), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=score-summary.xlsx; filename*=UTF-8''{quote(filename)}", "Cache-Control": "private, no-store"})
